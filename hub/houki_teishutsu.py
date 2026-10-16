"""houki_teishutsu — HOUKI-TEISHUTSU-1: 相続放棄 家裁提出パッケージ（提出前チェック・送付状 docx・提出日の記録）

QUEUE-2026-10-16-R F-3。feature/houki-e2e-1 の上。
正本: docs/souzoku-houki/09-shinjutsu-package.md（提出パッケージ）・06-documents（書類）・03-jukuryo-kikan（熟慮期間）。

 1) 提出前チェック（checklist）: 申述書・代理人目録・事情説明書（上申書）の添付、戸籍チェックリスト（App 40「書類チェック」）の充足、
    申述人情報の必須欄（hub/houki_shinjutsu.build_fill_data の拒否理由をそのまま使う）、熟慮期間の残日数（hub/houki_jukuryo.compute）
    を一覧にする。すべて ok のときだけ画面に「送付状を生成」「提出日を記録」を出す（**提出自体は人が行う**・表示のみ）。
 2) 送付状 docx: 雛形 docx_templates/houki/送付状.docx（scripts/make_soufujo_template.py で収載・全文明朝・太字なし・表／テキストボックス
    不使用の素の段落・各 {{ }} は単一 run・SHA pin）。差し込みは 宛先家裁・事件名・申述人・同封書類一覧・日付・事務所情報（OFFICE_* env）。
    同封書類一覧は **実際に添付されている生成済み書類**＋受領済みの戸籍等＋収入印紙・予納郵券（hub/houki_fees）から作る。
 3) 収入印紙・予納郵券は hub/houki_fees（定数 1 か所・要大野確認・未登録の家裁は「要確認」）。
 4) 提出日は人が /app/houki/submit で記録する（App 40「申述提出日」に $revision CAS で書く）。既存 HOUKI-JUKURYO は
    「申述提出日が空」の案件だけを対象にするため、記録すると 14／7 日前通知は止まる（本 module は通知ロジックに触れない）。

規律: 送信なし（LINE push ゼロ）・kintone 書込は 申述提出日 と FILE 欄「送付状」（あれば）だけ・RV-10（画面・ログに住所・電話・LINE userId を
出さない）・雛形改変は SHA pin で拒否。
"""

import datetime
import hashlib
import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from hub import houki_fees, houki_jukuryo
from hub import houki_dairinin as D
from hub import houki_koseki_checklist as K
from hub import houki_shinjutsu as S
from hub.docx_builder import fill_runs, to_wareki
from hub.redact import emit

logger = logging.getLogger("hub.houki_teishutsu")

TEMPLATE_PATH = "docx_templates/houki/送付状.docx"
TEMPLATE_SHA256 = "596b3ee15710d88666e792ca58101a027bfc42d02611fa6df638a2de32172c9e"
FONT_JA = "MS Mincho"
TITLE = "送　付　状"
KI = "記"
IJO = "以上"
BODY = "上記事件について、下記の書類を送付いたしますので、よろしくお取り計らいくださいますようお願い申し上げます。"
# 段落の並び（各 {{ }} は単一 run・表・テキストボックス・太字は使わない）。alignment: C=中央・R=右・L=左
LINES: tuple = (
    ("R", "{{日付}}"),
    ("L", "{{宛先家裁}}　御中"),
    ("L", ""),
    ("R", "{{事務所住所}}"),
    ("R", "{{事務所名}}"),
    ("R", "弁護士　{{代理人氏名}}"),
    ("R", "電話　{{電話}}　／　FAX　{{FAX}}"),
    ("L", ""),
    ("C", TITLE),
    ("L", ""),
    ("L", "事件名　{{事件名}}"),
    ("L", "申述人　{{申述人}}"),
    ("L", ""),
    ("L", BODY),
    ("L", ""),
    ("C", KI),
    ("L", ""),
    ("L", "{{同封書類}}"),
    ("L", ""),
    ("R", IJO),
)
PLACEHOLDERS = ("{{日付}}", "{{宛先家裁}}", "{{事務所住所}}", "{{事務所名}}", "{{代理人氏名}}", "{{電話}}", "{{FAX}}", "{{事件名}}",
                "{{申述人}}", "{{同封書類}}")

FIELD_SHINJUTSU = "申述書"
FIELD_DAIRININ = "代理人目録"
FIELD_JIJOU = "事情説明書"
FIELD_SOUFUJO = "送付状"                      # App 40 FILE（CU 追加前提・無ければ docx をダウンロードで返す）
FIELD_SUBMITTED = houki_jukuryo.FIELD_SUBMITTED   # 申述提出日
FIELD_COURT = "管轄家庭裁判所"
FIELD_CUSTOMER = "顧客名"
FIELD_DECEASED = "被相続人氏名"
REQUIRED_FIELDS = (FIELD_CUSTOMER, "住所", "生年月日", FIELD_DECEASED, FIELD_COURT)
DOC_LABELS = {FIELD_SHINJUTSU: "相続放棄申述書", FIELD_DAIRININ: "代理人目録", FIELD_JIJOU: "上申書（事情説明書）"}
_DOC_FIELDS = (FIELD_SHINJUTSU, FIELD_DAIRININ, FIELD_JIJOU)
_FW_DIGITS = str.maketrans("0123456789", "０１２３４５６７８９")
_JST = datetime.timezone(datetime.timedelta(hours=9))


class TeishutsuIntegrityError(RuntimeError):
    """雛形 pin 不一致・生成物の凍結検証失敗（添付しない）。"""


class TeishutsuRejection(RuntimeError):
    def __init__(self, reasons: list[str]):
        super().__init__(" / ".join(reasons))
        self.reasons = reasons


@dataclass(frozen=True)
class CheckItem:
    key: str
    label: str
    ok: bool
    detail: str


def _v(record: dict, code: str) -> str:
    return str(((record or {}).get(code) or {}).get("value") or "").strip()


def _attached(record: dict, code: str) -> bool:
    v = (record or {}).get(code)
    return bool(v and isinstance(v.get("value"), list) and v["value"])


def today_jst() -> datetime.date:
    return datetime.datetime.now(_JST).date()


# ── 1) 提出前チェック ────────────────────────────────────────────────────────
def _koseki_rows(record: dict) -> list:
    rows = K._rows(record)
    out = []
    for r in rows:
        name, target = K._cell(r, K.COL_NAME), K._cell(r, K.COL_TARGET)
        if name in K._KOSEKI_DOCS or (name == K.DOC_OTHER and "戸籍" in target):
            out.append(r)
    return out


def koseki_item(record: dict) -> CheckItem:
    rows = K._rows(record)
    if not rows:
        return CheckItem("koseki", "戸籍チェックリスト", False, "書類チェックに行がありません（初期行の生成待ち）")
    ks = _koseki_rows(record)
    states = [K._cell(r, K.COL_STATE) for r in ks]
    received = states.count(K.STATE_RECEIVED)
    not_needed = states.count(K.STATE_NOT_NEEDED)
    pending = len(states) - received - not_needed
    ok = bool(ks) and pending == 0 and received >= 1
    return CheckItem("koseki", "戸籍チェックリスト", ok, f"戸籍系 {len(ks)} 行: 受領 {received}／不要 {not_needed}／未了 {pending}")


def required_item(record: dict, today: datetime.date) -> CheckItem:
    missing = [f for f in REQUIRED_FIELDS if not _v(record, f)]
    reasons = []
    try:
        S.build_fill_data(record, today)
    except S.ShinjutsuRejection as e:
        reasons = list(e.reasons)
    except Exception:
        reasons = ["shinjutsu_check_failed"]
    ok = not missing and not reasons
    detail = "すべて入力済み" if ok else ("未入力: " + "・".join(missing) + ("　申述書の拒否理由: " + " / ".join(reasons) if reasons else ""))
    return CheckItem("required", "申述人情報の必須欄", ok, detail)


def jukuryo_item(record: dict, today: datetime.date) -> CheckItem:
    comp = houki_jukuryo.compute(record, today)
    if comp is None:
        return CheckItem("jukuryo", "熟慮期間", False, "起算日_確定が未入力（期限を算出できません）")
    ok = comp.remaining >= 0
    return CheckItem("jukuryo", "熟慮期間", ok,
                     f"期限 {comp.deadline.isoformat()}・残 {comp.remaining} 日" + ("" if ok else "（期限超過）"))


def checklist(record: dict, today: datetime.date | None = None) -> list[CheckItem]:
    today = today or today_jst()
    items = []
    for code in _DOC_FIELDS:
        items.append(CheckItem(code, DOC_LABELS[code], _attached(record, code), "添付あり" if _attached(record, code) else "未添付"))
    items.append(koseki_item(record))
    items.append(required_item(record, today))
    items.append(jukuryo_item(record, today))
    return items


def is_ready(items: list[CheckItem]) -> bool:
    return all(i.ok for i in items)


def submitted_on(record: dict) -> str:
    return _v(record, FIELD_SUBMITTED)


# ── 同封書類・差し込み ────────────────────────────────────────────────────────
def enclosures(record: dict) -> list[str]:
    """送付状の同封書類一覧（添付されている生成済み書類＋受領済みの戸籍等＋印紙・郵券）。"""
    out = [DOC_LABELS[c] for c in _DOC_FIELDS if _attached(record, c)]
    for r in K._rows(record):
        if K._cell(r, K.COL_STATE) != K.STATE_RECEIVED:
            continue
        name, target = K._cell(r, K.COL_NAME), K._cell(r, K.COL_TARGET)
        if name in (K.DOC_STAMP, K.DOC_POSTAGE) or not name:
            continue
        out.append(f"{name}（{target}）" if target and name != K.DOC_POA else name)
    out.append(f"収入印紙　{houki_fees.fmt_yen(houki_fees.stamp_yen(1))}分")
    p = houki_fees.postage_for(_v(record, FIELD_COURT))
    out.append(f"予納郵券　{houki_fees.fmt_yen(p.yen)}")
    return out


def enclosure_text(items: list[str]) -> str:
    return "\n".join(f"{str(i).translate(_FW_DIGITS)}　{name}" for i, name in enumerate(items, start=1))


def build_fill(record: dict, today: datetime.date | None = None) -> dict:
    today = today or today_jst()
    reasons = [f"missing_{f}" for f in (FIELD_COURT, FIELD_CUSTOMER, FIELD_DECEASED) if not _v(record, f)]
    if reasons:
        raise TeishutsuRejection(reasons)
    fill = {"{{日付}}": to_wareki(today), "{{宛先家裁}}": _v(record, FIELD_COURT), "{{事件名}}": D.case_name(record),
            "{{申述人}}": _v(record, FIELD_CUSTOMER), "{{同封書類}}": enclosure_text(enclosures(record))}
    office = D.office_values()
    for key in ("{{事務所住所}}", "{{事務所名}}", "{{代理人氏名}}", "{{電話}}", "{{FAX}}"):
        fill[key] = office.get(key, "")
    return fill


# ── 雛形・生成・凍結検証 ─────────────────────────────────────────────────────
def verify_template_bytes(docx_bytes: bytes) -> None:
    if hashlib.sha256(docx_bytes).hexdigest() != TEMPLATE_SHA256:
        raise TeishutsuIntegrityError("template hash mismatch")


def _expected_texts(filled: dict | None = None) -> list[str]:
    out = []
    for _al, text in LINES:
        t = text
        for k, v in (filled or {}).items():
            t = t.replace(k, v)
        out.append(t)
    return out


def verify_structure(doc, *, filled: dict | None = None) -> None:
    """段落の並びが LINES と一致・表／テキストボックスなし・太字なし・各プレースホルダは単一 run に 1 つ（雛形）／残存なし（生成物）。"""
    paras = list(doc.paragraphs)
    texts = [p.text for p in paras]
    if texts != _expected_texts(filled):
        raise TeishutsuIntegrityError("paragraph sequence mismatch")
    if doc.tables:
        raise TeishutsuIntegrityError("table present")
    xml = doc.element.xml
    if "txbx" in xml or "<w:tbl>" in xml or "<w:b/>" in xml or "<w:b " in xml:
        raise TeishutsuIntegrityError("textbox/table/bold present")
    for p in paras:
        for r in p.runs:
            if r.bold:
                raise TeishutsuIntegrityError("bold run")
            rf = r._element.rPr.rFonts if r._element.rPr is not None else None
            ea = rf.get(_qn("w:eastAsia")) if rf is not None else None
            if ea not in (None, FONT_JA):
                raise TeishutsuIntegrityError("font mismatch")
    normal = doc.styles["Normal"]
    ea = normal.element.rPr.rFonts.get(_qn("w:eastAsia")) if normal.element.rPr is not None and normal.element.rPr.rFonts is not None else None
    if ea != FONT_JA:
        raise TeishutsuIntegrityError("normal style font mismatch")
    if filled is None:
        for ph in PLACEHOLDERS:
            hits = [r for p in paras for r in p.runs if r.text == ph or ph in r.text]
            if len(hits) != 1 or sum(r.text.count(ph) for r in hits) != 1:
                raise TeishutsuIntegrityError("placeholder not single run")
    else:
        if any("{{" in t for t in texts):
            raise TeishutsuIntegrityError("placeholder remains")


def _qn(tag: str) -> str:
    from docx.oxml.ns import qn
    return qn(tag)


def render(fill: dict, template_path: str | Path = TEMPLATE_PATH) -> bytes:
    data = Path(template_path).read_bytes()
    verify_template_bytes(data)
    doc = Document(io.BytesIO(data))
    verify_structure(doc)
    for p in doc.paragraphs:
        fill_runs(p, fill)
    out = io.BytesIO()
    doc.save(out)
    docx_bytes = out.getvalue()
    verify_structure(Document(io.BytesIO(docx_bytes)), filled=fill)
    return docx_bytes


def generate(record: dict, today: datetime.date | None = None) -> bytes:
    """App 40 レコード → 送付状 docx（雛形 pin・構造・差し込み後の逐語一致を検証）。"""
    fill = build_fill(record, today)
    docx_bytes = render(fill)
    logger.info("[TEISHUTSU] soufujo generated bytes=%s", emit(len(docx_bytes), "count", "log", "operator"))
    return docx_bytes


def parse_date(text) -> datetime.date | None:
    s = str(text or "").strip()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return None
    try:
        return datetime.date.fromisoformat(s)
    except ValueError:
        return None
