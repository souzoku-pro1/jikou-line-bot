"""houki_jijou — HOUKI-JIJOU-1: 事情説明書（上申書）の LLM 下書き・素材一致検査・docx 生成（hub 側）

正本: docs/souzoku-houki/06-documents.md §3（事情説明書・上申書・3ヶ月超過型・Claude 下書きは時系列整理と文章化の補助・
主張の根拠は起算点メモ〔人の判断〕・AI が独自の法的主張を生成することを禁止・承認の二重化）・09 §1 ⑦。
QUEUE-2026-10-13-D D-5。司令塔仮置き:
- 項目（5 節）: 被相続人との関係／相続開始を知った経緯と日付／放棄する理由／債務の認識／遺産の処分の有無。
- 雛形: 申述書と同型の docx（docx_templates/houki/上申書_熟慮期間.docx・scripts/make_jijou_template.py が生成・SHA pin）。
- 下書きは App 40 の素材欄（SOURCE_FIELDS の閉集合）から生成し、生成物に素材に無い事実を入れない。ガード
  （check_against_source）: 下書き中の日付・金額・固有名詞が素材に無ければ違反（生成は止める・人が素材を直す）。
- 二重承認: PWA（/app/houki/jijou/{App 40 No.}）で大野が下書きを確認・修正し「承認して生成」→ docx 生成 → App 40 の
  FILE 欄「事情説明書」へ添付。送信なし。承認なしでは生成・添付しない（GET は下書き表示のみ・kintone 書込ゼロ）。
- RV-10: ログ・通知に氏名・住所・本文を載せない（レコード No.・件数・固定語彙のみ）。
"""

import datetime
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

from docx import Document

from hub.docx_builder import fill_runs
from hub.houki_shinjutsu import _split_court
from hub.houki_soufu_letter import wareki

TEMPLATE_PATH = "docx_templates/houki/上申書_熟慮期間.docx"
TEMPLATE_SHA256 = "0a05be024489974995407bb3731980df657e886eccb2e147296a7c2541c8e036"

TITLE = "上　申　書"
SUBTITLE = "（相続放棄の申述に関する事情説明書）"
HEAD_ROWS: tuple[tuple[str, str], ...] = (
    ("事件名", "{{事件名}}"),
    ("裁判所", "{{裁判所名}}"),
    ("申述人", "{{申述人氏名}}"),
    ("被相続人", "{{被相続人氏名}}"),
)
# 5 節（見出し・キー・プレースホルダ）。順序・逐語とも凍結
SECTIONS: tuple[tuple[str, str, str], ...] = (
    ("kankei", "１　被相続人との関係", "{{関係}}"),
    ("keii", "２　相続の開始を知った経緯と日付", "{{経緯}}"),
    ("riyu", "３　相続放棄をする理由", "{{理由}}"),
    ("saimu", "４　債務の認識", "{{債務}}"),
    ("shobun", "５　遺産の処分の有無", "{{処分}}"),
)
SECTION_KEYS = tuple(k for k, _, _ in SECTIONS)
FOOTER_ROWS: tuple[tuple[str, str], ...] = (
    ("作成日", "{{日付}}"),
    ("申述人代理人弁護士", "{{代理人氏名}}"),
)
LABELS = tuple(l for l, _ in HEAD_ROWS) + tuple(h for _, h, _ in SECTIONS) + tuple(l for l, _ in FOOTER_ROWS)
PLACEHOLDERS = tuple(p for _, p in HEAD_ROWS) + tuple(p for _, _, p in SECTIONS) + tuple(p for _, p in FOOTER_ROWS)
CLOSING = "以上のとおり、上申いたします。"

FIELD_ATTACHMENT = "事情説明書"              # App 40 FILE（大野が CU で追加・無い間は添付せず通知）
# 素材欄（閉集合・これ以外の事実は下書きに入れない）
SOURCE_FIELDS: tuple = (
    "顧客名", "被相続人氏名", "続柄", "続柄その他", "死亡日_申告", "死亡を知った日_申告", "相続人と知った日_申告",
    "日付申告メモ", "知った経緯", "財産_負債", "財産_現金預貯金", "財産_不動産", "財産_有価証券", "財産処分有無",
    "訴訟督促有無", "起算日_確定", "起算点メモ", "相続の開始を知った日", "死亡日", "管轄家庭裁判所",
)
MAX_SECTION_CHARS = 1200
MAX_DRAFT_TOKENS = 1500
LAWYER_CHECK = "【弁護士確認】"
_JST = datetime.timezone(datetime.timedelta(hours=9))

DRAFT_TOOL = {
    "name": "draft_jijou",
    "description": "相続放棄の上申書（事情説明書）の 5 節を、与えられた素材だけから下書きする",
    "input_schema": {
        "type": "object",
        "properties": {k: {"type": "string", "description": h} for k, h, _ in SECTIONS},
        "required": list(SECTION_KEYS),
    },
}
DRAFT_SYSTEM = (
    "あなたは法律事務所の文書作成補助です。家庭裁判所へ提出する相続放棄の上申書（事情説明書）の 5 節を、"
    "下に示す【素材】（kintone の案件記録）だけから下書きしてください。\n"
    "絶対規則:\n"
    "- 素材に無い事実（日付・金額・人名・会社名・出来事）を書かない。推測で補わない。不明な点は「【弁護士確認】」と書いて空ける。\n"
    "- 法的主張・結論（「熟慮期間は経過していない」等）を書かない。起算点に関する記述は素材の『起算点メモ』の内容を整理して書く。\n"
    "- 敬体・一人称は「申述人」。各節 400 字以内。箇条書き記号・markdown を使わない。\n"
    "- 日付は素材の表記（YYYY-MM-DD）をそのまま使う（和暦変換は後工程が行う）。\n"
)


class JijouIntegrityError(RuntimeError):
    """雛形 pin 不一致・生成物の凍結検証失敗（添付しない）。"""


class JijouGuardError(ValueError):
    """素材一致検査の違反（固定語彙の列）。"""

    def __init__(self, violations: list):
        super().__init__("jijou_guard")
        self.violations = list(violations)


def _fv(record: dict, code: str) -> str:
    return str(((record or {}).get(code) or {}).get("value") or "").strip()


def today_jst() -> datetime.date:
    return datetime.datetime.now(_JST).date()


def source_values(record: dict) -> dict:
    """素材欄の値（空は除く）。"""
    return {c: _fv(record, c) for c in SOURCE_FIELDS if _fv(record, c)}


def source_text(record: dict) -> str:
    return "\n".join(f"{k}: {v}" for k, v in source_values(record).items())


# ── 素材一致検査（純関数） ──────────────────────────────────────────────────────
_RE_DATE_ISO = re.compile(r"(\d{4})[-/年]\s*(\d{1,2})[-/月]\s*(\d{1,2})日?")
_RE_DATE_WAREKI = re.compile(r"(令和|平成|昭和)\s*(元|\d{1,2})年\s*(\d{1,2})月\s*(\d{1,2})日")
_RE_YM = re.compile(r"(\d{4})年\s*(\d{1,2})月(?!\s*\d)")
_RE_AMOUNT = re.compile(r"(\d[\d,]*)\s*(万円|円)")
_RE_PROPER = re.compile(r"[ァ-ヴー]{3,}|(?:[ァ-ヴー]{2,}|[一-龥]{2,}|[A-Za-z]{2,})(?:株式会社|有限会社|銀行|信用金庫|信用組合|債権回収|ファイナンス|カード|クレジット|ローン)|株式会社[ァ-ヴー一-龥A-Za-z]{1,20}")
_ERA = {"令和": 2018, "平成": 1988, "昭和": 1925}


def _norm(s: str) -> str:
    return unicodedata.normalize("NFKC", str(s or ""))


def _dates(text: str) -> set:
    out = set()
    t = _norm(text)
    for m in _RE_DATE_WAREKI.finditer(t):
        y = _ERA[m.group(1)] + (1 if m.group(2) == "元" else int(m.group(2)))
        out.add(f"{y:04d}-{int(m.group(3)):02d}-{int(m.group(4)):02d}")
    for m in _RE_DATE_ISO.finditer(t):
        out.add(f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}")
    return out


def _months(text: str) -> set:
    t = _norm(text)
    out = {f"{int(m.group(1)):04d}-{int(m.group(2)):02d}" for m in _RE_YM.finditer(t)}
    return out | {d[:7] for d in _dates(text)}


def _amounts(text: str) -> set:
    out = set()
    for m in _RE_AMOUNT.finditer(_norm(text)):
        n = int(m.group(1).replace(",", ""))
        out.add(n * 10000 if m.group(2) == "万円" else n)
    return out


def _proper_nouns(text: str) -> set:
    return {m.group(0) for m in _RE_PROPER.finditer(_norm(text))}


def check_against_source(sections: dict, record: dict) -> list[str]:
    """下書き（5 節）の日付・年月・金額・固有名詞が素材に無ければ違反（固定語彙 "節:種別:値"）。"""
    src = source_text(record)
    src_dates, src_months, src_amounts = _dates(src), _months(src), _amounts(src)
    src_norm = _norm(src)
    v: list[str] = []
    for key in SECTION_KEYS:
        text = str((sections or {}).get(key) or "")
        for d in sorted(_dates(text) - src_dates):
            v.append(f"{key}:date:{d}")
        for ym in sorted(_months(text) - src_months - {d[:7] for d in src_dates}):
            v.append(f"{key}:month:{ym}")
        for a in sorted(_amounts(text) - src_amounts):
            v.append(f"{key}:amount:{a}")
        for n in sorted(_proper_nouns(text)):
            if n not in src_norm:
                v.append(f"{key}:noun:{n}")
    return v


# ── 下書き（Claude・tool 強制） ────────────────────────────────────────────────
def parse_draft(tool_input) -> dict | None:
    if not isinstance(tool_input, dict):
        return None
    out = {}
    for k in SECTION_KEYS:
        val = tool_input.get(k)
        if not isinstance(val, str):
            return None
        out[k] = val.strip()[:MAX_SECTION_CHARS]
    return out


async def draft_sections(record: dict) -> dict:
    """Claude で 5 節を下書き（素材だけ・tool 強制）。解析不能は ValueError。"""
    import anthropic
    import chat_responder as cr
    from claude_gateway import create_message_with_fallback
    client = anthropic.AsyncAnthropic(api_key=cr._ANTHROPIC_KEY)
    resp = await create_message_with_fallback(
        client, context="相続放棄 上申書下書き", max_tokens=MAX_DRAFT_TOKENS,
        system=[{"type": "text", "text": DRAFT_SYSTEM}],
        tools=[json.loads(json.dumps(DRAFT_TOOL, ensure_ascii=False))],
        tool_choice={"type": "tool", "name": DRAFT_TOOL["name"]},
        messages=[{"role": "user", "content": "【素材】\n" + source_text(record)}])
    block = next((b for b in resp.content if getattr(b, "type", "") == "tool_use"), None)
    parsed = parse_draft(getattr(block, "input", None)) if block else None
    if parsed is None:
        raise ValueError("jijou_draft_unparseable")
    return parsed


# ── 差し込み・雛形検証・生成 ────────────────────────────────────────────────────
def case_name(record: dict) -> str:
    return f"被相続人　{_fv(record, '被相続人氏名')}　相続放棄申述事件"


def _court_name(record: dict) -> str:
    raw = _fv(record, "管轄家庭裁判所")
    if not raw:
        return ""
    before, after = _split_court(raw)
    return f"{before}家庭裁判所{after}" if before or after else raw


def build_fill(record: dict, sections: dict, *, today: datetime.date | None = None) -> dict:
    """差し込み値。節は本文そのまま（素材一致検査は呼出側が先に行う）。"""
    today = today or today_jst()
    fill = {
        "{{事件名}}": case_name(record),
        "{{裁判所名}}": _court_name(record),
        "{{申述人氏名}}": _fv(record, "顧客名"),
        "{{被相続人氏名}}": _fv(record, "被相続人氏名"),
        "{{日付}}": wareki(today),
        "{{代理人氏名}}": str(__import__("os").environ.get("OFFICE_ATTORNEY", "") or ""),
    }
    for key, _h, ph in SECTIONS:
        text = str((sections or {}).get(key) or "").strip()
        fill[ph] = text if text else LAWYER_CHECK
    return fill


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def verify_template_bytes(docx_bytes: bytes) -> None:
    if hashlib.sha256(docx_bytes).hexdigest() != TEMPLATE_SHA256:
        raise JijouIntegrityError("template hash mismatch")


def template_structure(doc) -> tuple:
    texts = [p.text for p in _all_paragraphs(doc)]
    return (texts[0] if texts else "", tuple(t for t in texts if t in LABELS), CLOSING in texts)


def verify_structure(doc, *, filled: dict | None = None) -> None:
    """見出し・ラベル・結語の逐語と、差し込み後のプレースホルダ残存 0 を検査。"""
    title, labels, closing = template_structure(doc)
    # python-docx の段落列挙は本文→表の順（文書順ではない）のため、ラベルは集合と個数で照合する
    if title != TITLE or sorted(labels) != sorted(LABELS) or not closing:
        raise JijouIntegrityError("template structure mismatch")
    joined = "\n".join(p.text for p in _all_paragraphs(doc))
    if filled is None:
        for ph in PLACEHOLDERS:
            if joined.count(ph) != 1:
                raise JijouIntegrityError("placeholder count mismatch")
    elif "{{" in joined or "}}" in joined:
        raise JijouIntegrityError("unfilled placeholder")


def render(fill: dict, template_path: str | Path = TEMPLATE_PATH) -> bytes:
    data = Path(template_path).read_bytes()
    verify_template_bytes(data)
    doc = Document(io.BytesIO(data))
    verify_structure(doc)
    for p in _all_paragraphs(doc):
        if "{{" in p.text:
            fill_runs(p, fill)
    verify_structure(doc, filled=fill)
    buf = io.BytesIO()
    doc.save(buf)
    out = buf.getvalue()
    verify_structure(Document(io.BytesIO(out)), filled=fill)
    return out


def generate(record: dict, sections: dict, *, today: datetime.date | None = None) -> bytes:
    """素材一致検査 → 差し込み → 凍結検証。違反は JijouGuardError（生成しない）。"""
    v = check_against_source(sections, record)
    if v:
        raise JijouGuardError(v)
    return render(build_fill(record, sections, today=today))
