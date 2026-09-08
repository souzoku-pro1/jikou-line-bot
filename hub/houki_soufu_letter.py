"""相続放棄 受理通知送付状（hub/houki_soufu_letter・HOUKI-SOUFU-1）

債権者宛の送付状 docx を、弁護士確定文（BODY_PARAGRAPHS・凍結）のテンプレート
docx_templates/houki/受理通知送付状.docx（scripts/make_houki_soufu_letter.py で生成・
SHA-256 pin）に App 40（相続放棄案件）と 債権者一覧 の行の値を差し込んで生成する。

- 差し込み 15 個（PLACEHOLDERS）。各キーは雛形で単一 run（fill_runs の契約）。
- 日付系（{{日付}} {{受理日}} {{被相続人生年月日}} {{死亡日}} {{申述人生年月日}}）は和暦
  （hub/docx_builder.to_wareki を流用）。App 40 の 生年月日・被相続人生年月日 は
  SINGLE_LINE_TEXT（実測 2026-09-08）のため、ISO 形式（YYYY-MM-DD）なら和暦へ変換し、
  それ以外の文字列は入力どおり差し込む（弁護士が和暦で入力している場合を壊さない）。
- 生成後に残存プレースホルダ 0 と本文段落の凍結 sha256 を検証する（fail-closed）。
"""

import hashlib
import io
import re
from datetime import date
from pathlib import Path

from docx import Document

from hub.docx_builder import fill_runs, to_wareki

TEMPLATE_PATH = "docx_templates/houki/受理通知送付状.docx"
# 生成物の SHA-256（scripts/make_houki_soufu_letter.py の出力を pin。雛形改変で生成拒否）
TEMPLATE_SHA256 = "e082fb5de9ad2f537b55e18bb68524ef9724b557acc069a880b7be7ca67e9383"

TITLE = "相続放棄申述受理のご通知兼書類送付のご案内"

# 弁護士確定文（凍結・改変禁止。1 行=1 段落・空行は空段落）
BODY_PARAGRAPHS = (
    "拝啓　時下ますますご清栄のこととお慶び申し上げます。",
    "当職は、{{申述人氏名}} 氏の代理人として、被相続人 {{被相続人氏名}} 氏に係る相続放棄の申述を行いました。",
    "同申述は、{{家庭裁判所名}}（事件番号：{{事件番号}}）において、{{受理日}} 付で受理されましたので、ご通知申し上げます。併せて、相続放棄申述受理通知書の写しを同封いたしますので、ご確認ください。",
    "つきましては、被相続人の債務に関し、上記申述人に対する請求・督促はお控えくださいますようお願い申し上げます。",
    "なお、本件に関するお問い合わせは、当職宛てにお願いいたします。",
    "敬具",
    "",
    "記",
    "",
    "【被相続人の表示】",
    "氏　　名　{{被相続人氏名}}",
    "最後の住所　{{被相続人最後の住所}}",
    "生年月日　{{被相続人生年月日}}",
    "死亡年月日　{{死亡日}}",
    "",
    "【申述人の表示】",
    "氏　　名　{{申述人氏名}}",
    "住　　所　{{申述人住所}}",
    "生年月日　{{申述人生年月日}}",
    "",
    "同封書類",
    "相続放棄申述受理通知書（写し）　1通",
    "以上",
)


def body_sha256(lines) -> str:
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


BODY_SHA256 = "652b789360689a26fbaeb60a0f453790d6035b28befb01a4f446cefba069ed32"   # テストで pin

PLACEHOLDERS = (
    "{{日付}}", "{{宛先郵便番号}}", "{{宛先住所}}", "{{宛先名}}", "{{事務所署名ブロック}}",
    "{{申述人氏名}}", "{{被相続人氏名}}", "{{家庭裁判所名}}", "{{事件番号}}", "{{受理日}}",
    "{{被相続人最後の住所}}", "{{被相続人生年月日}}", "{{死亡日}}", "{{申述人住所}}", "{{申述人生年月日}}",
)
DATE_PLACEHOLDERS = ("{{日付}}", "{{受理日}}", "{{被相続人生年月日}}", "{{死亡日}}", "{{申述人生年月日}}")
RECIPIENT_HONORIFIC = "御中"

# App 40 のフィールドコード → 差し込みキー（送付状の必須値）
CASE_FIELD_TO_KEY = {
    "顧客名": "{{申述人氏名}}",
    "住所": "{{申述人住所}}",
    "生年月日": "{{申述人生年月日}}",
    "被相続人氏名": "{{被相続人氏名}}",
    "被相続人最後の住所": "{{被相続人最後の住所}}",
    "被相続人生年月日": "{{被相続人生年月日}}",
    "死亡日": "{{死亡日}}",                  # 確定値（死亡日_申告 は使わない）
    "管轄家庭裁判所": "{{家庭裁判所名}}",
    "事件番号": "{{事件番号}}",
    "受理日": "{{受理日}}",
}
REQUIRED_CASE_FIELDS = tuple(CASE_FIELD_TO_KEY)

_ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class LetterIntegrityError(Exception):
    """雛形の改変・差し込み残存（fail-closed）。"""


def wareki(d: date) -> str:
    """令和/平成/昭和 X年X月X日（hub/docx_builder.to_wareki は昭和を持たないため、
    昭和期のみ本関数が補い、平成以降は to_wareki に委譲＝時効側の表記と同一）。"""
    if d >= date(1989, 1, 8):
        return to_wareki(d)
    if d >= date(1926, 12, 25):
        n = d.year - 1925
        return f"昭和{'元' if n == 1 else n}年{d.month}月{d.day}日"
    return d.strftime("%Y年%m月%d日")


def wareki_text(value: str) -> str:
    """ISO 日付は和暦へ。それ以外の文字列（弁護士入力の和暦等）はそのまま。"""
    s = str(value or "").strip()
    if _ISO_RE.match(s):
        try:
            return wareki(date.fromisoformat(s))
        except ValueError:
            return s
    return s


def _v(record: dict, code: str) -> str:
    return str((record.get(code) or {}).get("value") or "").strip()


def missing_case_fields(case: dict) -> list[str]:
    """送付状の必須値のうち空の欄コード（値は返さない）。"""
    return [code for code in REQUIRED_CASE_FIELDS if not _v(case, code)]


def build_letter_data(case: dict, creditor_name: str, creditor_zip: str, creditor_addr: str,
                      office_signature: str, today: date | None = None) -> dict:
    """15 個の差し込み値。日付系は和暦。"""
    today = today or date.today()
    data = {
        "{{日付}}": wareki(today),
        "{{宛先郵便番号}}": f"〒{creditor_zip.strip()}" if creditor_zip.strip() else "",
        "{{宛先住所}}": creditor_addr.strip(),
        "{{宛先名}}": f"{creditor_name.strip()}　{RECIPIENT_HONORIFIC}",
        "{{事務所署名ブロック}}": office_signature,
    }
    for code, key in CASE_FIELD_TO_KEY.items():
        val = _v(case, code)
        data[key] = wareki_text(val) if key in DATE_PLACEHOLDERS else val
    return data


def _fill_runs_multiline(para, mapping: dict) -> None:
    """fill_runs と同じ run 保持の差し込みに、値の改行を <w:br/> として展開する。"""
    for r in para.runs:
        if "{{" not in r.text:
            continue
        text = r.text
        for k, v in mapping.items():
            text = text.replace(k, v)
        if "\n" not in text:
            r.text = text
            continue
        lines = text.split("\n")
        r.text = lines[0]
        for line in lines[1:]:
            r.add_break()
            r.add_text(line)


def verify_template_bytes(docx_bytes: bytes) -> None:
    if hashlib.sha256(docx_bytes).hexdigest() != TEMPLATE_SHA256:
        raise LetterIntegrityError("template sha mismatch")


def template_body(docx_bytes: bytes) -> list[str]:
    """雛形の本文段落（表題の 2 つ後から末尾）。"""
    doc = Document(io.BytesIO(docx_bytes))
    texts = [p.text for p in doc.paragraphs]
    i = texts.index(TITLE)
    return texts[i + 2:]


def render_letter(data: dict, template_path: str | Path = TEMPLATE_PATH) -> bytes:
    """雛形 SHA 検証 → run 保持で差し込み → 残存プレースホルダ 0 を検証。"""
    raw = Path(template_path).read_bytes()
    verify_template_bytes(raw)
    if template_body(raw) != list(BODY_PARAGRAPHS) or body_sha256(BODY_PARAGRAPHS) != BODY_SHA256:
        raise LetterIntegrityError("frozen body mismatch")
    missing = [k for k in PLACEHOLDERS if k not in data]
    if missing:
        raise LetterIntegrityError("data keys missing")
    doc = Document(io.BytesIO(raw))
    for para in doc.paragraphs:
        if para.runs and any("{{" in r.text for r in para.runs):
            fill_runs(para, {k: v for k, v in data.items() if "\n" not in v})
            _fill_runs_multiline(para, {k: v for k, v in data.items() if "\n" in v})
    if any("{{" in p.text for p in doc.paragraphs):
        raise LetterIntegrityError("placeholder remains")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
