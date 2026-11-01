"""jikou_receipt — 領収書（xlsx）の採番・金額計算・生成・保管（JIKOU-RECEIPT-1・QUEUE-2026-10-20-HI I-2）

- 起点: F-2 の「入金確認」（App 21 入金状況=入金済み）。生成は人の明示操作（/app/jikou/receipt の確認付き POST）のみ。
- 金額: 料金規則（chat_responder.fee_total_yen＝1〜2 社目 44,000 円・3 社目以降 22,000 円・税込）から取り、手入力させない。
  内訳は税抜単価 UNIT_BASE_EXCL=40,000／UNIT_EXTRA_EXCL=20,000・消費税 TAX_RATE_PERCENT=10%・端数切捨て（**要大野確認**）。
  追加依頼は差額のみ（既発行の社数を超えた分だけ・3 社目以降の単価で）。
- 番号: 年（JST）＋連番（receipt_sequence 表・UNIQUE(business, year, seq)・衝突は再試行）。再発行は同番号＋「再発行」表示。
- 雛形: docx_templates/jikou/領収書.xlsx（{{ }} 差し込み・SHA pin）。正本は大野作成＝本票は仮雛形（scripts/make_jikou_receipt_template.py）。
- 但し書きは hub/jikou_doc_texts RC-1（要承認）。承認台帳（hub/approved_texts）に本文 SHA が載るまで発行不可（text_unapproved）。
- PDF: xlsx→PDF の既存変換経路は無い（contract_pdf は docx 専用）→ xlsx のまま保管し保留（票の指示どおり）。
- 登録番号欄は env OFFICE_INVOICE_NO があるときだけ表示。数字は半角・桁区切り（number_format '#,##0'）。
- RV-10: 本文に入る相談者情報は宛名（顧客名）のみ。ログ・通知は案件 No.・番号・金額のみ。
"""

from __future__ import annotations

import datetime
import hashlib
import io
import logging
import os
import uuid
from dataclasses import asdict, dataclass

import sqlalchemy as sa
from sqlalchemy.exc import IntegrityError

from hub import approved_texts, docx_builder
from hub import jikou_doc_texts as T
from hub import jikou_followup as F
from hub.db import session_scope
from hub.redact import emit

logger = logging.getLogger("hub.jikou_receipt")

BUSINESS = "jikou"
TABLE_NAME = "jikou_receipt"
SEQ_TABLE_NAME = "receipt_sequence"
TABLE_NAMES = (TABLE_NAME, SEQ_TABLE_NAME)
TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                             "docx_templates", "jikou", "領収書.xlsx")
TEMPLATE_SHA256 = "ddb55904f42183e429dc39122620d41d2f5e770fb5ff62aa8ea4ff2ec16ee9d7"       # scripts/make_jikou_receipt_template.py の出力（仮雛形）

# 料金規則（税抜内訳・要大野確認）。税込の合計は chat_responder.fee_total_yen と一致すること（test で pin）
UNIT_BASE_EXCL = 40_000        # 1〜2 社目の 1 社あたり（税抜）
UNIT_EXTRA_EXCL = 20_000       # 3 社目以降の 1 社あたり（税抜）
BASE_MAX_COUNT = 2
TAX_RATE_PERCENT = 10          # 消費税率（%）・端数切捨て
INVOICE_NO_ENV = "OFFICE_INVOICE_NO"
OFFICE_ENV = {"事務所名": "OFFICE_NAME", "事務所住所": "OFFICE_ADDRESS", "電話": "OFFICE_TEL"}
FIELD_PAYMENT = "入金状況"
PAYMENT_PAID = "入金済み"
NAME_FIELD = "顧客名"
TEXT_KEY_NOTE = "RC-1"
REISSUE_LABEL = "再発行"
_ALLOC_RETRIES = 20
_JST = datetime.timezone(datetime.timedelta(hours=9))

metadata = sa.MetaData()
receipt_sequence = sa.Table(
    SEQ_TABLE_NAME, metadata,
    sa.Column("business", sa.Text, primary_key=True),
    sa.Column("year", sa.Integer, primary_key=True),
    sa.Column("seq", sa.Integer, primary_key=True),
    sa.Column("allocated_at", sa.DateTime(timezone=True), nullable=False),
)
jikou_receipt = sa.Table(
    TABLE_NAME, metadata,
    sa.Column("receipt_id", sa.Text, primary_key=True),
    sa.Column("business", sa.Text, nullable=False),
    sa.Column("case_record_id", sa.Text, nullable=False),
    sa.Column("number", sa.Text, nullable=False),              # 例 2026-0001（再発行は同番号）
    sa.Column("year", sa.Integer, nullable=False),
    sa.Column("seq", sa.Integer, nullable=False),
    sa.Column("issued_on", sa.Date, nullable=False),
    sa.Column("addressee", sa.Text, nullable=False),
    sa.Column("creditor_names", sa.Text, nullable=False),
    sa.Column("count_total", sa.Integer, nullable=False),      # 発行時点の通算社数
    sa.Column("count_billed", sa.Integer, nullable=False),     # 本領収書で請求した社数（差額）
    sa.Column("qty1", sa.Integer, nullable=False), sa.Column("unit1", sa.Integer, nullable=False), sa.Column("amt1", sa.Integer, nullable=False),
    sa.Column("qty2", sa.Integer, nullable=False), sa.Column("unit2", sa.Integer, nullable=False), sa.Column("amt2", sa.Integer, nullable=False),
    sa.Column("subtotal", sa.Integer, nullable=False), sa.Column("tax", sa.Integer, nullable=False), sa.Column("total", sa.Integer, nullable=False),
    sa.Column("reissue", sa.Boolean, nullable=False),
    sa.Column("reissue_of", sa.Text, nullable=True),
    sa.Column("xlsx", sa.LargeBinary, nullable=False),
    sa.Column("sha256", sa.Text, nullable=False),
    sa.Column("texts_version", sa.Text, nullable=False),
    sa.Column("template_sha256", sa.Text, nullable=False),
    sa.Column("created_by", sa.Text, nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.CheckConstraint("count_billed >= 1", name="ck_jikou_receipt_billed"),
    sa.Index("ix_jikou_receipt_case", "business", "case_record_id"),
    sa.Index("ix_jikou_receipt_number", "business", "number"),
)


class IssueRefused(RuntimeError):
    """固定コード: not_paid / no_creditors / nothing_to_bill / name_missing / text_unapproved / texts_pin / template_sha /
    placeholder_left / number_alloc / not_found"""


@dataclass(frozen=True)
class Breakdown:
    count_total: int
    count_billed: int
    qty1: int
    unit1: int
    amt1: int
    qty2: int
    unit2: int
    amt2: int
    subtotal: int
    tax: int
    total: int


def compute(count_total: int, already_billed: int = 0) -> Breakdown:
    """通算 count_total 社のうち、既発行 already_billed 社を超えた分だけを請求する内訳（差額）。"""
    total = max(int(count_total), 0)
    billed = min(max(int(already_billed), 0), total)
    positions = range(billed + 1, total + 1)
    qty1 = sum(1 for p in positions if p <= BASE_MAX_COUNT)
    qty2 = sum(1 for p in positions if p > BASE_MAX_COUNT)
    amt1, amt2 = qty1 * UNIT_BASE_EXCL, qty2 * UNIT_EXTRA_EXCL
    subtotal = amt1 + amt2
    tax = subtotal * TAX_RATE_PERCENT // 100          # 端数切捨て（要大野確認）
    return Breakdown(total, total - billed, qty1, UNIT_BASE_EXCL, amt1, qty2, UNIT_EXTRA_EXCL, amt2, subtotal, tax, subtotal + tax)


def consistent_with_fee_rule(max_count: int = 6) -> bool:
    """税込合計が料金規則（chat_responder.fee_total_yen）と一致するか（起動時は呼ばない・test で pin）。"""
    import chat_responder
    return all(compute(n).total == chat_responder.fee_total_yen(n) for n in range(0, max_count + 1))


def template_ok() -> bool:
    try:
        with open(TEMPLATE_PATH, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest() == TEMPLATE_SHA256
    except OSError:
        return False


def _fv(record: dict | None, code: str) -> str:
    return str(((record or {}).get(code) or {}).get("value") or "").strip()


def _now() -> datetime.datetime:
    return datetime.datetime.now(datetime.timezone.utc)


def today_jst(now: datetime.datetime | None = None) -> datetime.date:
    return (now or _now()).astimezone(_JST).date()


def format_number(year: int, seq: int) -> str:
    return f"{int(year)}-{int(seq):04d}"


async def allocate_number(year: int, *, now: datetime.datetime | None = None) -> tuple[int, str]:
    """年ごとの連番を 1 つ確保する（UNIQUE 衝突は再試行・重複なし）。戻り値 (seq, number)。"""
    now = now or _now()
    for _ in range(_ALLOC_RETRIES):
        try:
            async with session_scope() as s:
                cur = (await s.execute(sa.select(sa.func.max(receipt_sequence.c.seq)).where(
                    receipt_sequence.c.business == BUSINESS, receipt_sequence.c.year == int(year)))).scalar()
                seq = int(cur or 0) + 1
                await s.execute(sa.insert(receipt_sequence).values(business=BUSINESS, year=int(year), seq=seq, allocated_at=now))
            return seq, format_number(year, seq)
        except IntegrityError:
            continue
    raise IssueRefused("number_alloc")


def _note_text(creditor_names: str) -> str:
    body = approved_texts.approved_text(TEXT_KEY_NOTE, T.text(TEXT_KEY_NOTE))
    if body is None:
        raise IssueRefused("text_unapproved")
    return body.replace("{{債権者名}}", creditor_names)


def _office_values() -> dict:
    vals = {"{{" + k + "}}": os.environ.get(env, "").strip() for k, env in OFFICE_ENV.items()}
    inv = os.environ.get(INVOICE_NO_ENV, "").strip()
    vals["{{登録番号}}"] = f"登録番号　{inv}" if inv else ""        # 無いときは欄ごと空
    return vals


def build_xlsx(values: dict, amounts: dict) -> bytes:
    """雛形の {{ }} を置換する。amounts は数値セル（int・number_format '#,##0'）。残った {{ は拒否。"""
    from openpyxl import load_workbook
    if not template_ok():
        raise IssueRefused("template_sha")
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or "{{" not in v:
                continue
            hit = next((k for k in amounts if v == k), None)
            if hit is not None:
                cell.value = int(amounts[hit])
                cell.number_format = "#,##0"
                continue
            for k, rep in values.items():
                v = v.replace(k, str(rep))
            cell.value = v
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                raise IssueRefused("placeholder_left")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render(number: str, issued_on: datetime.date, addressee: str, creditor_names: str, b: Breakdown, *, reissue: bool) -> bytes:
    values = {"{{領収書番号}}": number, "{{領収日}}": docx_builder.to_wareki(issued_on), "{{宛名}}": addressee,
              "{{債権者名}}": creditor_names, "{{但し書き}}": _note_text(creditor_names),
              "{{再発行}}": REISSUE_LABEL if reissue else "", **_office_values()}
    amounts = {"{{合計}}": b.total, "{{数量1}}": b.qty1, "{{単価1}}": b.unit1, "{{金額1}}": b.amt1,
               "{{数量2}}": b.qty2, "{{単価2}}": b.unit2, "{{金額2}}": b.amt2, "{{小計}}": b.subtotal, "{{消費税}}": b.tax}
    return build_xlsx(values, amounts)


async def list_receipts(case_record_id: str) -> list[dict]:
    async with session_scope() as s:
        t = jikou_receipt
        rows = (await s.execute(sa.select(t.c.receipt_id, t.c.number, t.c.issued_on, t.c.count_total, t.c.count_billed, t.c.total,
                                          t.c.reissue, t.c.reissue_of, t.c.created_at)
                                .where(t.c.business == BUSINESS, t.c.case_record_id == str(case_record_id))
                                .order_by(t.c.created_at.asc()))).fetchall()
    return [dict(r._mapping) for r in rows]


async def get_receipt(case_record_id: str, receipt_id: str) -> dict | None:
    async with session_scope() as s:
        t = jikou_receipt
        r = (await s.execute(sa.select(t).where(t.c.business == BUSINESS, t.c.case_record_id == str(case_record_id),
                                                t.c.receipt_id == str(receipt_id)))).first()
    return dict(r._mapping) if r else None


async def billed_count(case_record_id: str) -> int:
    """既発行（再発行を除く）の請求済み社数の合計。"""
    async with session_scope() as s:
        t = jikou_receipt
        v = (await s.execute(sa.select(sa.func.coalesce(sa.func.sum(t.c.count_billed), 0)).where(
            t.c.business == BUSINESS, t.c.case_record_id == str(case_record_id), t.c.reissue.is_(False)))).scalar()
    return int(v or 0)


async def _store(case_record_id: str, *, number: str, year: int, seq: int, issued_on: datetime.date, addressee: str,
                 creditor_names: str, b: Breakdown, reissue: bool, reissue_of: str | None, xlsx: bytes, actor: str,
                 now: datetime.datetime) -> dict:
    rid = uuid.uuid4().hex
    async with session_scope() as s:
        await s.execute(sa.insert(jikou_receipt).values(
            receipt_id=rid, business=BUSINESS, case_record_id=str(case_record_id), number=number, year=year, seq=seq,
            issued_on=issued_on, addressee=addressee, creditor_names=creditor_names, count_total=b.count_total,
            count_billed=b.count_billed, qty1=b.qty1, unit1=b.unit1, amt1=b.amt1, qty2=b.qty2, unit2=b.unit2, amt2=b.amt2,
            subtotal=b.subtotal, tax=b.tax, total=b.total, reissue=reissue, reissue_of=reissue_of, xlsx=xlsx,
            sha256=hashlib.sha256(xlsx).hexdigest(), texts_version=T.TABLE_SHA256_PIN[:16], template_sha256=TEMPLATE_SHA256,
            created_by=actor, created_at=now))
    logger.info("[JIKOU_RECEIPT] stored case=%s number=%s total=%s reissue=%s",
                emit(str(case_record_id), "record_id", "log", "operator"), emit(number, "record_id", "log", "operator"),
                emit(b.total, "count", "log", "operator"), emit(int(reissue), "count", "log", "operator"))
    return {"receipt_id": rid, "number": number, **asdict(b), "reissue": reissue}


async def issue(case_record_id: str, *, actor: str, record: dict | None = None, now: datetime.datetime | None = None) -> dict:
    """入金済みの案件に対し、未請求の社数分（差額）の領収書を新番号で発行する。"""
    now = now or _now()
    if not T.verify_pin():
        raise IssueRefused("texts_pin")
    if record is None:
        from hub import kintone
        try:
            record = await kintone.get_record(F.APP21, str(case_record_id))
        except Exception as e:  # noqa: BLE001
            raise IssueRefused("name_missing") from e
    if _fv(record, FIELD_PAYMENT) != PAYMENT_PAID:
        raise IssueRefused("not_paid")
    creditors = F.creditors_of(record)
    if not creditors:
        raise IssueRefused("no_creditors")
    addressee = _fv(record, NAME_FIELD)
    if not addressee:
        raise IssueRefused("name_missing")
    already = await billed_count(case_record_id)
    b = compute(len(creditors), already)
    if b.count_billed <= 0:
        raise IssueRefused("nothing_to_bill")
    names = "・".join(n for _, n in creditors[already:])      # 本領収書で請求する社のみ
    issued_on = today_jst(now)
    xlsx_probe = _render("0000-0000", issued_on, addressee, names, b, reissue=False)   # 文言・雛形の検査を採番より先に
    del xlsx_probe
    seq, number = await allocate_number(issued_on.year, now=now)
    xlsx = _render(number, issued_on, addressee, names, b, reissue=False)
    return await _store(case_record_id, number=number, year=issued_on.year, seq=seq, issued_on=issued_on, addressee=addressee,
                        creditor_names=names, b=b, reissue=False, reissue_of=None, xlsx=xlsx, actor=actor, now=now)


async def reissue(case_record_id: str, receipt_id: str, *, actor: str, now: datetime.datetime | None = None) -> dict:
    """同じ番号・同じ金額で「再発行」表示付きの領収書を作る（元の行は残す）。"""
    now = now or _now()
    orig = await get_receipt(case_record_id, receipt_id)
    if orig is None:
        raise IssueRefused("not_found")
    b = Breakdown(orig["count_total"], orig["count_billed"], orig["qty1"], orig["unit1"], orig["amt1"], orig["qty2"],
                  orig["unit2"], orig["amt2"], orig["subtotal"], orig["tax"], orig["total"])
    xlsx = _render(orig["number"], orig["issued_on"], orig["addressee"], orig["creditor_names"], b, reissue=True)
    return await _store(case_record_id, number=orig["number"], year=orig["year"], seq=orig["seq"], issued_on=orig["issued_on"],
                        addressee=orig["addressee"], creditor_names=orig["creditor_names"], b=b, reissue=True,
                        reissue_of=str(receipt_id), xlsx=xlsx, actor=actor, now=now)
