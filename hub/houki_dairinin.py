"""代理人目録の生成（HOUKI-DAIRININ-1・hub 側モジュール）

相続放棄申述に添える「代理人目録」docx を、テンプレ docx_templates/houki/代理人目録.docx
（scripts/make_dairinin_template.py が python-docx で生成・体裁は申述書テンプレの用紙/余白/既定フォントを複製・
SHA-256 pin）に App 40（相続放棄案件）と事務所情報（env OFFICE_*）を差し込んで生成する。
QUEUE-2026-10-12-C C-3。司令塔仮置き:
- 項目: 事件名・申述人氏名・住所・本籍・代理人（氏名・事務所名・住所・電話・FAX・弁護士会・登録番号）・裁判所名・
  事件番号（空欄可）・日付。
- 事務所情報は env OFFICE_* を正とし、弁護士会・登録番号は新 env OFFICE_BAR・OFFICE_REG（値は大野が Railway に設定）。
  未設定の項目は空欄で生成し、受け口が業務 LINE に警告を 1 回出す（生成は止めない）。
- FAX は env OFFICE_FAX の値のみを使う（8374/8370 のどちらが正か不明のため雛形に固定文を持たない。既存雛形
  〔契約書・通知書〕は触らない）。
- 申述人本籍は App 40 に欄が無い（HEARING-FIX1 実測）ため空欄（手書き運用）。
- 生成の拒否条件は申述書と同じ 10 種（hub/houki_shinjutsu.build_fill_data を検証にだけ流用＝ShinjutsuRejection
  をそのまま送出）。承認ゲート: 生成は自動・添付のみ（送信なし）。

凍結検証（受理通知送付状と同型）: テンプレの SHA-256 と 見出し/ラベルの逐語（LABELS）を照合し、差し込み後に
プレースホルダ残存 0・ラベル不変を検査する（DairininIntegrityError）。
RV-10: ログ・通知に氏名・住所を載せない（レコード No.・欄名のみ）。
"""

import datetime
import hashlib
import io
import os
from pathlib import Path

from docx import Document

from hub import houki_shinjutsu
from hub.docx_builder import fill_runs
from hub.houki_soufu_letter import wareki

TEMPLATE_PATH = "docx_templates/houki/代理人目録.docx"
TEMPLATE_SHA256 = "0a73c6876b2f74ff74dd64ef0e47d866df07f2dd91def3fa214fc90b5d14d77d"      # scripts/make_dairinin_template.py の出力を pin（雛形改変で生成拒否）

TITLE = "代　理　人　目　録"
# 表の行（ラベル → プレースホルダ）。順序・逐語とも凍結（雛形と本定数の両方を検証する）
ROWS: tuple[tuple[str, str], ...] = (
    ("事件名", "{{事件名}}"),
    ("裁判所", "{{裁判所名}}"),
    ("事件番号", "{{事件番号}}"),
    ("申述人　氏名", "{{申述人氏名}}"),
    ("申述人　住所", "{{申述人住所}}"),
    ("申述人　本籍", "{{申述人本籍}}"),
    ("代理人　氏名", "{{代理人氏名}}"),
    ("代理人　事務所名", "{{事務所名}}"),
    ("代理人　住所", "{{事務所住所}}"),
    ("代理人　電話", "{{電話}}"),
    ("代理人　FAX", "{{FAX}}"),
    ("代理人　弁護士会", "{{弁護士会}}"),
    ("代理人　登録番号", "{{登録番号}}"),
    ("作成日", "{{日付}}"),
)
LABELS = tuple(label for label, _ in ROWS)
PLACEHOLDERS = tuple(ph for _, ph in ROWS)
FOOTER = "上記のとおり代理人を定めました。"

# 事務所情報（env・OFFICE_BAR / OFFICE_REG は本票で追加）→ プレースホルダ
OFFICE_ENV_TO_KEY = (
    ("OFFICE_ATTORNEY", "{{代理人氏名}}"),
    ("OFFICE_NAME", "{{事務所名}}"),
    ("OFFICE_ADDRESS", "{{事務所住所}}"),
    ("OFFICE_TEL", "{{電話}}"),
    ("OFFICE_FAX", "{{FAX}}"),
    ("OFFICE_BAR", "{{弁護士会}}"),
    ("OFFICE_REG", "{{登録番号}}"),
)
OFFICE_ENVS = tuple(env for env, _ in OFFICE_ENV_TO_KEY)
_JST = datetime.timezone(datetime.timedelta(hours=9))


class DairininIntegrityError(RuntimeError):
    """テンプレ pin 不一致・生成物の凍結検証失敗（添付しない）。"""


def _fv(record: dict, code: str) -> str:
    return str(((record or {}).get(code) or {}).get("value") or "").strip()


def today_jst() -> datetime.date:
    return datetime.datetime.now(_JST).date()


def office_values() -> dict:
    """env から事務所情報（未設定は ""）。住所は OFFICE_ZIP があれば「〒zip 住所」。"""
    out = {key: os.environ.get(env, "").strip() for env, key in OFFICE_ENV_TO_KEY}
    zip_ = os.environ.get("OFFICE_ZIP", "").strip()
    if zip_ and out["{{事務所住所}}"]:
        out["{{事務所住所}}"] = f"〒{zip_}　{out['{{事務所住所}}']}"
    return out


def missing_office_envs() -> list[str]:
    """未設定の env 名（値は返さない・通知の文面に使う）。"""
    return [env for env in OFFICE_ENVS if not os.environ.get(env, "").strip()]


def case_name(record: dict) -> str:
    deceased = _fv(record, "被相続人氏名")
    return f"被相続人　{deceased}　相続放棄申述事件" if deceased else ""


def build_fill(record: dict, today: datetime.date | None = None) -> dict:
    """App 40 レコード → 14 個の差し込み値。拒否条件は申述書と同じ（ShinjutsuRejection をそのまま送出）。"""
    houki_shinjutsu.build_fill_data(record, today=today)      # 検証にだけ使う（10 種の拒否理由）
    today = today or today_jst()
    fill = {
        "{{事件名}}": case_name(record),
        "{{裁判所名}}": _fv(record, "管轄家庭裁判所"),
        "{{事件番号}}": _fv(record, "事件番号"),            # 空欄可
        "{{申述人氏名}}": _fv(record, "顧客名"),
        "{{申述人住所}}": _fv(record, "住所"),
        "{{申述人本籍}}": _fv(record, "申述人本籍"),        # App 40 に欄なし＝通常は空欄
        "{{日付}}": wareki(today),
    }
    fill.update(office_values())
    return fill


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def verify_template_bytes(docx_bytes: bytes) -> None:
    if hashlib.sha256(docx_bytes).hexdigest() != TEMPLATE_SHA256:
        raise DairininIntegrityError("template sha mismatch")


def template_structure(doc) -> tuple[str, tuple, str]:
    """(表題, ((ラベル, 値セル本文), ...), 末尾文)。"""
    paras = [p.text for p in doc.paragraphs]
    title = paras[0] if paras else ""
    rows = tuple((t.rows[i].cells[0].text, t.rows[i].cells[1].text)
                 for t in doc.tables[:1] for i in range(len(t.rows)))
    footer = next((x for x in paras[1:] if x.strip()), "")
    return title, rows, footer


def verify_structure(doc, *, filled: dict | None = None) -> None:
    """表題・ラベル・末尾文の逐語と、値セル（差し込み前=プレースホルダ／後=値）を検査。"""
    title, rows, footer = template_structure(doc)
    if title != TITLE or footer != FOOTER or len(rows) != len(ROWS):
        raise DairininIntegrityError("frozen text mismatch")
    for (label, ph), (got_label, got_value) in zip(ROWS, rows):
        if got_label != label:
            raise DairininIntegrityError("label mismatch")
        expected = ph if filled is None else filled.get(ph, "")
        if got_value != expected:
            raise DairininIntegrityError("value mismatch")
    if filled is not None and any("{{" in p.text for p in _all_paragraphs(doc)):
        raise DairininIntegrityError("placeholder remains")


def render(fill: dict, template_path: str | Path = TEMPLATE_PATH) -> bytes:
    raw = Path(template_path).read_bytes()
    verify_template_bytes(raw)
    doc = Document(io.BytesIO(raw))
    verify_structure(doc)
    missing = [k for k in PLACEHOLDERS if k not in fill]
    if missing:
        raise DairininIntegrityError("data keys missing")
    for p in _all_paragraphs(doc):
        if any("{{" in r.text for r in p.runs):
            fill_runs(p, fill)
    verify_structure(doc, filled=fill)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def generate(record: dict, today: datetime.date | None = None) -> bytes:
    """検証済み代理人目録 docx（拒否は ShinjutsuRejection・完全性は DairininIntegrityError）。
    kintone・通知は受け口側の責務。"""
    return render(build_fill(record, today=today))
