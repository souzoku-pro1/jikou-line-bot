"""HOUKI-SHOKUGYO-1（QUEUE-2026-10-13-D D-8）: App 40「職業」欄の接続（書込の可用性ゲート・申述書差し込み）

固定する仕様:
- HEARING_WRITABLE_FIELDS に 職業 を復帰。upsert 直前に hub/houki_shokugyo.available（form fields・1 時間キャッシュ・失敗は
  無い側）で検査し、無ければ 職業 だけ落として他欄を書く（会話を止めない＝HEARING-FIX1 の止血は維持）。
- 申述書: 職業 が非空なら「職業」ラベル右のセルに差し込む（雛形 SHA・canonical pin は不変・期待列と完全一致）。
  空なら従来と同一の生成物。セル構造の想定崩れは IntegrityError。
- ヒアリング 7 通（prompt）は不変。sink 違反ゼロ・allowlist 追加ゼロ。
"""

import asyncio
import hashlib
import io
import os
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document

for _k, _v in {
    "KINTONE_SUBDOMAIN": "testsub", "LINE_CHANNEL_SECRET": "dummy_secret",
    "LINE_CHANNEL_ACCESS_TOKEN": "dummy_token", "ANTHROPIC_API_KEY": "dummy_key",
    "KINTONE_APP_ID": "21", "KINTONE_API_TOKEN": "dummy",
    "SOUZOKU_KINTONE_APP_ID": "26", "SOUZOKU_KINTONE_API_TOKEN": "dummy",
    "APP_APPROVAL": "29", "TOKEN_APPROVAL": "d", "APP_HOUKI": "40", "TOKEN_HOUKI": "t40",
}.items():
    os.environ.setdefault(_k, _v)

from hub import houki_case_store as store  # noqa: E402
from hub import houki_shinjutsu as sj  # noqa: E402
from hub import houki_shokugyo as sg  # noqa: E402
from hub import kintone  # noqa: E402
from test_houki_shinjutsu import TODAY as _TODAY, _full_record  # noqa: E402
from test_sink_ast_policy import scan_source  # noqa: E402

REPO = Path(__file__).parent
HEARING_PROMPT_SHA256 = "12a1a94c85897d8e1ef747d08b673287e242760616643a0ea5b41e57d9465716"


def _run(coro):
    return asyncio.run(coro)


class TestAvailability(unittest.TestCase):
    def setUp(self):
        sg.clear_cache()

    def tearDown(self):
        sg.clear_cache()

    def test_available_cached_and_fail_closed(self):
        ff = AsyncMock(return_value={"職業": {"type": "SINGLE_LINE_TEXT"}})
        with patch.object(kintone, "get_form_fields", ff):
            self.assertTrue(_run(sg.available(now=0.0)))
            self.assertTrue(_run(sg.available(now=10.0)))
            self.assertEqual(ff.await_count, 1)                               # キャッシュ
            self.assertTrue(_run(sg.available(now=sg.CACHE_SECONDS + 1)))
            self.assertEqual(ff.await_count, 2)
        sg.clear_cache()
        with patch.object(kintone, "get_form_fields", AsyncMock(return_value={})):
            self.assertFalse(_run(sg.available()))
        sg.clear_cache()
        with patch.object(kintone, "get_form_fields", AsyncMock(side_effect=RuntimeError("boom"))):
            self.assertFalse(_run(sg.available()))
            self.assertFalse(_run(sg.available()))                            # 失敗はキャッシュしない

    def test_strip_if_unavailable(self):
        with patch.object(sg, "available", AsyncMock(return_value=False)):
            self.assertEqual(_run(sg.strip_if_unavailable({"顧客名": "x", "職業": "会社員"})), {"顧客名": "x"})
            self.assertEqual(_run(sg.strip_if_unavailable({"顧客名": "x"})), {"顧客名": "x"})
        with patch.object(sg, "available", AsyncMock(return_value=True)):
            self.assertEqual(_run(sg.strip_if_unavailable({"職業": "会社員"})), {"職業": "会社員"})
        self.assertEqual(sg.value_for_docx({"職業": {"value": " 会社\n員 " + "x" * 30}}), ("会社 員 " + "x" * 30)[:sg.MAX_CHARS])
        self.assertEqual(sg.value_for_docx({}), "")

    def test_writable_set_and_upsert_gate(self):
        self.assertIn("職業", store.HEARING_WRITABLE_FIELDS)
        out, _p, _c = store.split_valid_fields({"職業": "会社員", "顧客名": "山田"})
        self.assertEqual(out, {"職業": "会社員", "顧客名": "山田"})
        create = AsyncMock(return_value="77")
        with patch.object(kintone, "create_record", create), patch.object(sg, "available", AsyncMock(return_value=False)):
            rid = _run(store.upsert_case_fields("U" + "a" * 32, {"職業": "会社員", "顧客名": "山田"}, None))
        self.assertEqual(rid, "77")
        self.assertNotIn("職業", create.await_args.args[1])
        self.assertEqual(create.await_args.args[1]["顧客名"], "山田")
        update = AsyncMock()
        existing = {"$id": {"value": "77"}, "$revision": {"value": "3"}, "顧客名": {"value": "山田"}, "職業": {"value": ""}}
        with patch.object(kintone, "update_record", update), patch.object(sg, "available", AsyncMock(return_value=True)):
            _run(store.upsert_case_fields("U" + "a" * 32, {"職業": "会社員", "顧客名": "山田"}, existing))
        self.assertEqual(update.await_args.args[2], {"職業": "会社員"})


class TestShinjutsuCell(unittest.TestCase):
    def test_empty_value_keeps_output_identical_and_pins_unchanged(self):
        rec = _full_record()
        a = sj.generate(rec, today=_TODAY)
        rec2 = _full_record(職業="")
        b = sj.generate(rec2, today=_TODAY)
        self.assertEqual([p.text for p in sj._walk_paragraphs(Document(io.BytesIO(a)))],
                         [p.text for p in sj._walk_paragraphs(Document(io.BytesIO(b)))])
        data = Path(sj.TEMPLATE_PATH).read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(), sj.TEMPLATE_SHA256)
        self.assertEqual(sj.canonical_sha256(data), sj.TEMPLATE_CANONICAL_SHA256)

    def test_value_is_placed_in_shokugyo_cell_and_verified(self):
        rec = _full_record(職業="会社員")
        out = sj.generate(rec, today=_TODAY)
        doc = Document(io.BytesIO(out))
        ti, ri, ci = sj._SHOKUGYO_CELL
        from docx.table import _Cell
        cell = _Cell(doc.tables[ti].rows[ri]._tr.tc_lst[ci], doc.tables[ti])
        self.assertEqual([p.text for p in cell.paragraphs], ["", "会社員", "", ""])
        label = _Cell(doc.tables[ti].rows[ri]._tr.tc_lst[ci - 1], doc.tables[ti])
        self.assertIn("職業", label.text)
        # 期待列に無い差し込み（値の食い違い）は検証で拒否される
        fill = sj.build_fill_data(rec, today=_TODAY)
        circles, _ = sj._circles(rec)
        with self.assertRaises(sj.ShinjutsuIntegrityError):
            sj.verify_shinjutsu_docx(out, fill, circles, "自営業")
        with self.assertRaises(sj.ShinjutsuIntegrityError):
            sj.verify_shinjutsu_docx(out, fill, circles, "")
        # セル構造の想定崩れは IntegrityError
        tdoc = Document(sj.TEMPLATE_PATH)
        sj._fill_shokugyo_cell(tdoc, "x")
        with self.assertRaises(sj.ShinjutsuIntegrityError):
            sj._fill_shokugyo_cell(tdoc, "y")


class TestPins(unittest.TestCase):
    def test_sink_allowlist_and_prompt_unchanged(self):
        allow = (REPO / "redaction_sink_allowlist.json").read_text(encoding="utf-8")
        for name in ("hub/houki_shokugyo.py", "hub/houki_case_store.py", "hub/houki_shinjutsu.py"):
            self.assertEqual(scan_source((REPO / name).read_text(encoding="utf-8"), name), [], name)
        self.assertNotIn("houki_shokugyo", allow)
        src = (REPO / "hub" / "houki_case_store.py").read_text(encoding="utf-8")
        self.assertNotIn("実欄が存在しない（form fields", src)                   # stale コメントの更新
        self.assertIn("HOUKI-SHOKUGYO-1", src)
        from hub import houki_profile
        self.assertEqual(hashlib.sha256(houki_profile.HOUKI_HEARING_PROMPT.encode("utf-8")).hexdigest(),
                         HEARING_PROMPT_SHA256)                                      # ヒアリング 7 通（prompt）不変


if __name__ == "__main__":
    unittest.main()
