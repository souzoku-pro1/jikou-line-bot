"""HOUKI-DAIRININ-1（QUEUE-2026-10-12-C C-3）: 代理人目録 docx の生成・添付

固定する仕様:
- テンプレ docx_templates/houki/代理人目録.docx は SHA-256 pin・表題/ラベル/末尾文の逐語凍結・各プレースホルダは
  単一 run（fill_runs の契約）。改変は生成拒否（DairininIntegrityError）。
- 差し込み 14 個: 事件名（被相続人氏名から合成）・裁判所名・事件番号（空欄可）・申述人氏名/住所/本籍（本籍は App 40 に
  欄が無く空欄）・代理人 7 項目（env OFFICE_ATTORNEY/NAME/ADDRESS(+ZIP)/TEL/FAX/BAR/REG）・作成日（和暦）。
- 拒否条件は申述書と同じ 10 種（houki_shinjutsu.build_fill_data の ShinjutsuRejection をそのまま）。
- env 未設定は空欄で生成し、受け口が業務 LINE に警告 1 回（throttle key）。
- 受け口 /souzoku-houki/dairinin/{secret}: 認証は申述書と同じ。添付欄「代理人目録」が無ければ生成せず skip+通知 1 回。
  冪等（非空=skip）・CAS（409=cas_lost）・添付失敗=通知+500。送信はしない。
"""

import asyncio
import datetime
import hashlib
import io
import os
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

_ENV = {
    "ANTHROPIC_API_KEY": "dummy", "LINE_CHANNEL_SECRET": "dummy_secret",
    "LINE_CHANNEL_ACCESS_TOKEN": "dummy_token", "KINTONE_SUBDOMAIN": "testsub",
    "KINTONE_APP_ID": "21", "KINTONE_API_TOKEN": "dummy",
    "SOUZOKU_KINTONE_APP_ID": "26", "SOUZOKU_KINTONE_API_TOKEN": "dummy",
    "CLOUDSIGN_CLIENT_ID": "c", "CLOUDSIGN_WEBHOOK_SECRET": "cs",
    "KINTONE_WEBHOOK_TOKEN": "kintone-token", "DOCUMENT_WEBHOOK_SECRET": "doc-secret",
    "APP_APPROVAL": "29", "TOKEN_APPROVAL": "d", "HEALTHCHECK_DISABLED": "1",
    "STRIPE_WEBHOOK_SECRET": "w", "GOOGLE_VISION_API_KEY": "dummy_vision",
    "APP_CHATLOG": "28", "TOKEN_CHATLOG": "d", "APP_HOUKI": "40", "TOKEN_HOUKI": "d",
    "HOUKI_LINE_CHANNEL_SECRET": "houki_secret", "HOUKI_LINE_CHANNEL_ACCESS_TOKEN": "houki_token",
    "HOUKI_WEBHOOK_TOKEN": "houki-hook",
}
for _k, _v in _ENV.items():
    os.environ.setdefault(_k, _v)

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import dairinin_webhook as dw  # noqa: E402
import main  # noqa: E402
from hub import houki_dairinin as hd  # noqa: E402
from hub import houki_shinjutsu as hs  # noqa: E402
from hub.kintone import KintoneConflict, KintoneError  # noqa: E402
from test_houki_shinjutsu import _full_record  # noqa: E402
from test_sink_ast_policy import scan_source  # noqa: E402

REPO = Path(__file__).parent
_client = TestClient(main.app)
_URL = "/souzoku-houki/dairinin/houki-hook"
TODAY = datetime.date(2026, 10, 12)
OFFICE = {"OFFICE_ATTORNEY": "大野 太郎", "OFFICE_NAME": "大野法律事務所", "OFFICE_ZIP": "332-0031",
          "OFFICE_ADDRESS": "埼玉県川口市西青木2-1-45 レクイアーレ101号", "OFFICE_TEL": "048-299-2704",
          "OFFICE_FAX": "048-000-0000", "OFFICE_BAR": "埼玉弁護士会", "OFFICE_REG": "12345"}
_NO_OFFICE = {k: "" for k in OFFICE}


def _texts(docx_bytes: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            out.extend(c.text for c in row.cells)
    return out


# ── 1. テンプレ pin・構造 ──────────────────────────────────────────────────────
class TestTemplate(unittest.TestCase):
    def test_sha_pin_and_frozen_structure(self):
        raw = (REPO / hd.TEMPLATE_PATH).read_bytes()
        self.assertEqual(hashlib.sha256(raw).hexdigest(), hd.TEMPLATE_SHA256)
        doc = Document(io.BytesIO(raw))
        hd.verify_structure(doc)
        for _label, ph in hd.ROWS:
            runs = [r for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs
                    for r in p.runs if ph in r.text]
            self.assertEqual(len(runs), 1, ph)
            self.assertEqual(runs[0].text, ph)
        self.assertEqual(len(hd.ROWS), 14)
        self.assertEqual(hd.OFFICE_ENVS, ("OFFICE_ATTORNEY", "OFFICE_NAME", "OFFICE_ADDRESS", "OFFICE_TEL",
                                          "OFFICE_FAX", "OFFICE_BAR", "OFFICE_REG"))

    def test_page_setup_matches_shinjutsu_template(self):
        a = Document(str(REPO / hd.TEMPLATE_PATH)).sections[0]
        b = Document(str(REPO / hs.TEMPLATE_PATH)).sections[0]
        self.assertEqual((a.page_width, a.page_height, a.left_margin, a.right_margin, a.top_margin, a.bottom_margin),
                         (b.page_width, b.page_height, b.left_margin, b.right_margin, b.top_margin, b.bottom_margin))

    def test_existing_templates_untouched(self):
        # 既存雛形（契約書・受理通知送付状・申述書）には触らない（FAX の正本は保留一覧へ）
        from hub import houki_soufu_letter
        self.assertEqual(hashlib.sha256((REPO / houki_soufu_letter.TEMPLATE_PATH).read_bytes()).hexdigest(),
                         houki_soufu_letter.TEMPLATE_SHA256)
        self.assertEqual(hashlib.sha256((REPO / hs.TEMPLATE_PATH).read_bytes()).hexdigest(), hs.TEMPLATE_SHA256)


# ── 2. 差し込み値・拒否・env ─────────────────────────────────────────────────
class TestBuildFill(unittest.TestCase):
    def test_values_with_office_env(self):
        with patch.dict(os.environ, OFFICE):
            fill = hd.build_fill(_full_record(事件番号="令和8年（家）第123号"), today=TODAY)
        self.assertEqual(set(fill), set(hd.PLACEHOLDERS))
        self.assertEqual(fill["{{事件名}}"], "被相続人　山田太郎　相続放棄申述事件")
        self.assertEqual(fill["{{申述人氏名}}"], "山田花子")
        self.assertEqual(fill["{{事件番号}}"], "令和8年（家）第123号")
        self.assertEqual(fill["{{申述人本籍}}"], "")
        self.assertEqual(fill["{{日付}}"], "令和8年10月12日")
        self.assertEqual(fill["{{事務所住所}}"], "〒332-0031　埼玉県川口市西青木2-1-45 レクイアーレ101号")
        self.assertEqual((fill["{{弁護士会}}"], fill["{{登録番号}}"], fill["{{FAX}}"]),
                         ("埼玉弁護士会", "12345", "048-000-0000"))
        self.assertEqual(hd.missing_office_envs(), [] if all(os.environ.get(k) for k in hd.OFFICE_ENVS)
                         else hd.missing_office_envs())

    def test_missing_env_is_blank_and_listed(self):
        with patch.dict(os.environ, _NO_OFFICE):
            fill = hd.build_fill(_full_record(), today=TODAY)
            self.assertEqual([fill[k] for _e, k in hd.OFFICE_ENV_TO_KEY], [""] * 7)
            self.assertEqual(hd.missing_office_envs(), list(hd.OFFICE_ENVS))
            self.assertEqual(fill["{{事件番号}}"], "")                   # 空欄可
        with patch.dict(os.environ, {**OFFICE, "OFFICE_BAR": "", "OFFICE_ZIP": ""}):
            self.assertEqual(hd.missing_office_envs(), ["OFFICE_BAR"])
            self.assertEqual(hd.office_values()["{{事務所住所}}"], OFFICE["OFFICE_ADDRESS"])

    def test_rejections_are_the_same_ten_as_shinjutsu(self):
        with self.assertRaises(hs.ShinjutsuRejection) as cm:
            hd.build_fill(_full_record(顧客名=""), today=TODAY)
        self.assertIn(hs.REJECT_MISSING_NAME, cm.exception.reasons)
        with self.assertRaises(hs.ShinjutsuRejection) as cm:
            hd.build_fill(_full_record(管轄家庭裁判所="さいたま地方裁判所"), today=TODAY)
        self.assertIn(hs.REJECT_COURT_UNMAPPED, cm.exception.reasons)
        with self.assertRaises(hs.ShinjutsuRejection) as cm:
            hd.build_fill(_full_record(被相続人氏名="", 死亡日_申告=""), today=TODAY)
        self.assertEqual(set(cm.exception.reasons) & {hs.REJECT_MISSING_DECEASED, hs.REJECT_MISSING_DEATH},
                         {hs.REJECT_MISSING_DECEASED, hs.REJECT_MISSING_DEATH})


# ── 3. 生成と凍結検証 ─────────────────────────────────────────────────────────
class TestRender(unittest.TestCase):
    def test_generate_fills_all_and_keeps_labels(self):
        with patch.dict(os.environ, OFFICE):
            out = hd.generate(_full_record(), today=TODAY)
            fill = hd.build_fill(_full_record(), today=TODAY)
        texts = _texts(out)
        joined = "\n".join(texts)
        self.assertNotIn("{{", joined)
        self.assertIn("被相続人　山田太郎　相続放棄申述事件", texts)
        self.assertIn("埼玉弁護士会", texts)
        self.assertIn("令和8年10月12日", texts)
        for label in hd.LABELS:
            self.assertIn(label, texts)
        doc = Document(io.BytesIO(out))
        hd.verify_structure(doc, filled=fill)

    def test_tampered_template_is_refused(self):
        raw = (REPO / hd.TEMPLATE_PATH).read_bytes()
        with patch.object(hd, "TEMPLATE_SHA256", "0" * 64):
            with self.assertRaises(hd.DairininIntegrityError):
                hd.render({k: "" for k in hd.PLACEHOLDERS})
        doc = Document(io.BytesIO(raw))
        doc.tables[0].rows[0].cells[0].paragraphs[0].runs[0].text = "改変"
        with self.assertRaises(hd.DairininIntegrityError):
            hd.verify_structure(doc)
        with self.assertRaises(hd.DairininIntegrityError):
            hd.render({k: "" for k in hd.PLACEHOLDERS[:-1]})


# ── 4. 受け口 ───────────────────────────────────────────────────────────────
class _WebhookBase(unittest.TestCase):
    def setUp(self):
        dw.reset_field_cache()
        self.get = AsyncMock(return_value=_full_record())
        self.upload = AsyncMock(return_value="fk-1")
        self.update = AsyncMock()
        self.fields = AsyncMock(return_value={dw.FIELD_ATTACHMENT: {"type": "FILE"}})
        self.notify = AsyncMock(return_value=True)
        self._env = patch.dict(os.environ, OFFICE)
        self._env.start()
        for p in (patch.object(dw.hub_kintone, "get_record", self.get),
                  patch.object(dw.hub_kintone, "upload_file", self.upload),
                  patch.object(dw.hub_kintone, "update_record", self.update),
                  patch.object(dw.hub_kintone, "get_form_fields", self.fields),
                  patch("hub.notify.notify_admin_line", self.notify)):
            p.start()
            self.addCleanup(p.stop)
        self.addCleanup(self._env.stop)

    def post(self, url=_URL, body=None):
        return _client.post(url, json=body if body is not None else {"record_id": "12"})


class TestWebhookAuth(_WebhookBase):
    def test_auth_and_shape(self):
        with patch.dict(os.environ, {"HOUKI_WEBHOOK_TOKEN": ""}):
            self.assertEqual(self.post().status_code, 404)
        with patch.dict(os.environ, {"HOUKI_WEBHOOK_TOKEN": os.environ["KINTONE_WEBHOOK_TOKEN"]}):
            self.assertEqual(self.post("/souzoku-houki/dairinin/kintone-token").status_code, 404)
        self.assertEqual(self.post("/souzoku-houki/dairinin/wrong").status_code, 403)
        self.assertEqual(self.post(body={"app": {"id": "21"}, "record": {"$id": {"value": "12"}}}).json().get("skip"),
                         "app_mismatch")
        self.assertEqual(self.post(body={"x": 1}).json().get("skip"), "no_record_id")
        self.get.assert_not_awaited()
        self.upload.assert_not_awaited()


class TestWebhookFlow(_WebhookBase):
    def test_happy_path_attaches_once_without_notice(self):
        r = self.post(body={"app": {"id": "40"}, "record": {"$id": {"value": "12"}}})
        self.assertEqual((r.status_code, r.json().get("record_id")), (200, "12"))
        self.upload.assert_awaited_once()
        self.assertEqual(self.upload.await_args.args[1], "代理人目録_12_山田花子.docx")
        self.assertNotIn("{{", "\n".join(_texts(self.upload.await_args.args[2])))
        self.assertEqual(self.update.await_args.args[2], {dw.FIELD_ATTACHMENT: [{"fileKey": "fk-1"}]})
        self.assertEqual(self.update.await_args.kwargs.get("revision"), "5")
        self.notify.assert_not_awaited()

    def test_field_missing_skips_with_one_notice(self):
        self.fields.return_value = {"申述書": {"type": "FILE"}}
        self.assertEqual(self.post().json().get("skip"), "field_missing")
        self.assertEqual(self.post().json().get("skip"), "field_missing")
        self.get.assert_not_awaited()
        self.assertEqual(self.fields.await_count, 1)
        self.assertEqual(self.notify.await_count, 2)                        # throttle は notify 側（key を渡す）
        self.assertEqual(self.notify.await_args.kwargs.get("throttle_key"), "dairinin_field_missing")
        self.assertIn("代理人目録", self.notify.await_args.args[0])

    def test_office_env_missing_generates_blank_and_warns_with_key(self):
        with patch.dict(os.environ, {"OFFICE_BAR": "", "OFFICE_REG": ""}):
            r = self.post()
        self.assertEqual(r.json().get("record_id"), "12")
        self.upload.assert_awaited_once()
        self.assertEqual(self.notify.await_count, 1)
        text = self.notify.await_args.args[0]
        self.assertIn("OFFICE_BAR", text)
        self.assertIn("OFFICE_REG", text)
        self.assertNotIn("山田", text)
        self.assertEqual(self.notify.await_args.kwargs.get("throttle_key"), "dairinin_office_env")

    def test_already_attached_cas_lost_and_attach_failure(self):
        self.get.return_value = _full_record(代理人目録=[{"fileKey": "existing"}])
        self.assertEqual(self.post().json().get("skip"), "already_attached")
        self.upload.assert_not_awaited()
        self.get.return_value = _full_record()
        self.update.side_effect = KintoneConflict(409, "GAIA_CO02", "conflict")
        self.assertEqual(self.post().json().get("skip"), "cas_lost")
        self.update.side_effect = KintoneError(500, "GAIA_XX", "down")
        r = self.post()
        self.assertEqual(r.status_code, 500)
        self.assertIn("要確認", self.notify.await_args.args[0])
        self.assertNotIn("山田", self.notify.await_args.args[0])

    def test_rejection_notifies_closed_vocabulary(self):
        self.get.return_value = _full_record(顧客名="")
        r = self.post()
        self.assertEqual(r.json().get("skip"), "rejected")
        self.assertIn(hs.REJECT_MISSING_NAME, r.json().get("reasons"))
        self.upload.assert_not_awaited()
        self.assertIn("No.12", self.notify.await_args.args[0])


# ── 5. pin ─────────────────────────────────────────────────────────────────────
class TestPins(unittest.TestCase):
    def test_wiring_and_sink(self):
        for name in ("hub/houki_dairinin.py", "dairinin_webhook.py", "scripts/make_dairinin_template.py"):
            src = (REPO / name).read_text(encoding="utf-8")
            self.assertEqual(scan_source(src, name), [], name)
        allow = (REPO / "redaction_sink_allowlist.json").read_text(encoding="utf-8")
        self.assertNotIn("dairinin", allow)
        src = (REPO / "main.py").read_text(encoding="utf-8").split("\n")
        self.assertIn("app.include_router(dairinin_router)", "\n".join(src[-12:]))


if __name__ == "__main__":
    unittest.main()
