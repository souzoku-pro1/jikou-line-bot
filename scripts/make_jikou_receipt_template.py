# -*- coding: utf-8 -*-
"""領収書の **仮雛形** xlsx を生成する — JIKOU-RECEIPT-1（QUEUE-2026-10-20-HI I-2）

正本は大野作成の docx_templates/jikou/領収書.xlsx（{{ }} 差し込み欄あり）。本票時点で repo に無いため、票の差し込み欄と同じ配置で
仮雛形を作って進める（保留一覧に記載・正本が届いたら本スクリプトの出力を差し替え TEMPLATE_SHA256 を更新）。
差し込み欄: 領収日・領収書番号・宛名・債権者名・但し書き・数量1/単価1/金額1（1〜2 社目）・数量2/単価2/金額2（3 社目以降）・
小計・消費税・合計・再発行表示・事務所名/住所/電話・登録番号（env OFFICE_INVOICE_NO があるときだけ表示）。
セル式は持たない（金額はサーバ側 int 計算・ZAISAN-GEN-1 と同じ）。zip 時刻と docProps の日時を固定した決定的バイト列。
使い方: python scripts/make_jikou_receipt_template.py [<出力先>]
sink 規律: print は定数＋emit のみ。
"""
import datetime
import hashlib
import io
import os
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from hub.redact import emit  # noqa: E402

_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_DT = datetime.datetime(1980, 1, 1, 0, 0, 0)
FONT_JA = "ＭＳ 明朝"
OUT_DEFAULT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docx_templates", "jikou", "領収書.xlsx")

# セル配置（票の差し込み欄）。値はプレースホルダのみ（文言は hub/jikou_doc_texts・金額はサーバ側で書く）
CELLS = {
    "B2": "領 収 書",
    "E3": "No. {{領収書番号}}",
    "E4": "{{領収日}}",
    "B6": "{{宛名}}　様",
    "B8": "金額",
    "D8": "{{合計}}",
    "B9": "但し　{{但し書き}}",
    "B11": "内容", "C11": "数量", "D11": "単価", "E11": "金額",
    "B12": "時効援用手続き（1〜2社目）", "C12": "{{数量1}}", "D12": "{{単価1}}", "E12": "{{金額1}}",
    "B13": "時効援用手続き（3社目以降）", "C13": "{{数量2}}", "D13": "{{単価2}}", "E13": "{{金額2}}",
    "D14": "小計", "E14": "{{小計}}",
    "D15": "消費税（10%）", "E15": "{{消費税}}",
    "D16": "合計", "E16": "{{合計}}",
    "B18": "{{再発行}}",
    "B20": "{{事務所名}}",
    "B21": "{{事務所住所}}",
    "B22": "電話　{{電話}}",
    "B23": "{{登録番号}}",
    "B25": "対象債権者　{{債権者名}}",
}
AMOUNT_CELLS = ("D8", "C12", "D12", "E12", "C13", "D13", "E13", "E14", "E15", "E16")
PLACEHOLDERS = tuple(sorted({v[v.index("{{"):v.index("}}") + 2] for v in CELLS.values() if "{{" in v}))


def build_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "領収書"
    thin = Side(style="thin", color="000000")
    for ref, value in CELLS.items():
        c = ws[ref]
        c.value = value
        c.font = Font(name=FONT_JA, size=16 if ref == "B2" else 11, bold=False)
        if ref in AMOUNT_CELLS or ref.startswith("E"):
            c.alignment = Alignment(horizontal="right")
    for r in range(11, 17):
        for col in "BCDE":
            ws[f"{col}{r}"].border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18
    wb.properties.created = _FIXED_DT
    wb.properties.modified = _FIXED_DT
    wb.properties.creator = "jikou-line-bot"
    wb.properties.lastModifiedBy = "jikou-line-bot"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def deterministic_zip(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, src.read(info.filename))
    return buf.getvalue()


def main(argv) -> int:
    out = argv[1] if len(argv) > 1 else OUT_DEFAULT
    data = deterministic_zip(build_xlsx())
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "wb") as f:
        f.write(data)
    print("written: docx_templates/jikou/領収書.xlsx (provisional; or argv[1])")
    print("TEMPLATE_SHA256 =", emit(hashlib.sha256(data).hexdigest(), "record_id", "log", "operator"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
