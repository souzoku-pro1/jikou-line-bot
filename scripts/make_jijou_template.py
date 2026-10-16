"""上申書（事情説明書）雛形の収載スクリプト（HOUKI-JIJOU-1）

申述書テンプレ docx_templates/houki/相続放棄申述書.docx の用紙・余白・既定書式（sectPr・styles）を複製し、本文を
表題 + 頭書き（2 列の表）+ 5 節（見出し段落+本文プレースホルダ）+ 結語 + 末尾（2 列の表）に置き換えて
docx_templates/houki/上申書_熟慮期間.docx を生成する。各 {{ }} は単一 run。生成物は決定的 zip（タイムスタンプ固定）で、
SHA-256 を hub/houki_jijou.TEMPLATE_SHA256 に pin する（代理人目録・受理通知送付状と同型）。

実行: python scripts/make_jijou_template.py
"""

import hashlib
import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from hub.houki_jijou import (CLOSING, FOOTER_ROWS, HEAD_ROWS, SECTIONS, SUBTITLE,  # noqa: E402
                             TEMPLATE_SHA256, TITLE, verify_structure)

SOURCE = REPO / "docx_templates" / "houki" / "相続放棄申述書.docx"
OUT = REPO / "docx_templates" / "houki" / "上申書_熟慮期間.docx"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


def deterministic_zip(docx_bytes: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0
            dst.writestr(zi, src.read(info.filename))
    return out.getvalue()


def _table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, ph) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.paragraphs[0].add_run(label)
        c1.paragraphs[0].add_run(ph)
    return table


def build() -> bytes:
    doc = Document(str(SOURCE))
    body = doc.element.body
    for child in list(body.iterchildren()):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(SUBTITLE)
    doc.add_paragraph("")
    _table(doc, HEAD_ROWS)
    doc.add_paragraph("")
    for _key, heading, ph in SECTIONS:
        h = doc.add_paragraph()
        hr = h.add_run(heading)
        hr.bold = True
        doc.add_paragraph().add_run(ph)
        doc.add_paragraph("")
    doc.add_paragraph().add_run(CLOSING)
    doc.add_paragraph("")
    _table(doc, FOOTER_ROWS)
    verify_structure(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return deterministic_zip(buf.getvalue())


def main() -> int:
    data = build()
    OUT.write_bytes(data)
    digest = hashlib.sha256(data).hexdigest()
    if digest != TEMPLATE_SHA256:
        sys.stdout.write("template written; hub/houki_jijou.TEMPLATE_SHA256 must be updated to the new digest\n")
        return 1
    sys.stdout.write("template written; digest matches hub/houki_jijou.TEMPLATE_SHA256\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
