"""相続放棄 受理通知送付状テンプレートの収載スクリプト（HOUKI-SOUFU-1）

時効側 docx_templates/jikou/送付案内.docx（事務所正式書式）の書式（用紙・余白・
フォント・日付/宛先/署名ブロックの配置・表題の書式）を複製し、本文を弁護士確定文
（hub/houki_soufu_letter.BODY_PARAGRAPHS・凍結）に置き換えて
docx_templates/houki/受理通知送付状.docx を生成する。表は使わない。各 {{ }} は
単一 run。生成物は決定的 zip（タイムスタンプ固定）で、SHA-256 を
hub/houki_soufu_letter.TEMPLATE_SHA256 に pin する。

併せて Desktop\\claude\\受理通知送付状_v1.docx（同一バイト列）と、Word が使える環境では
受理通知送付状_v1.pdf（プレースホルダのままの見本）を出力する。

実行: python scripts/make_houki_soufu_letter.py
"""

import copy
import hashlib
import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from hub.houki_soufu_letter import (BODY_PARAGRAPHS, BODY_SHA256,  # noqa: E402
                                    PLACEHOLDERS, TITLE, body_sha256)

SOURCE = REPO / "docx_templates" / "jikou" / "送付案内.docx"
SOURCE_SHA256 = "5cca0d724de56577c2a12e606f595528b3bb3160c5ace9f85cd417a41c15e972"
OUT = REPO / "docx_templates" / "houki" / "受理通知送付状.docx"
DESKTOP = Path.home() / "Desktop" / "claude"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)

# 時効書式の段落番地（scratchpad 実測 2026-09-08）
_P_DATE, _P_ADDR, _P_NAME, _P_BLANK1, _P_SIGN, _P_BLANK2, _P_TITLE = 0, 1, 2, 3, 4, 5, 6
_P_BODY_NORMAL, _P_KEIGU, _P_CENTER, _P_IJO = 8, 14, 15, 18

RIGHT_ALIGNED = {"敬具", "以上"}
CENTERED = {"記"}


def deterministic_zip(docx_bytes: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0
            dst.writestr(zi, src.read(info.filename))
    return out.getvalue()


def die(msg: str) -> None:
    raise SystemExit(f"収載中止: {msg}")


def _set_single_run(para, text: str) -> None:
    """段落の先頭 run にテキストを置き、他 run を除去（プレースホルダは単一 run）。"""
    runs = para.runs
    if not runs:
        para.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _clone_para(src_para, text: str, after):
    """src_para の pPr/先頭 run の rPr を複製した新段落を after の直後に置く。"""
    new_p = copy.deepcopy(src_para._p)
    for r in list(new_p.iterchildren()):
        if r.tag.endswith("}r"):
            new_p.remove(r)
    after.addnext(new_p)
    from docx.text.paragraph import Paragraph
    para = Paragraph(new_p, src_para._parent)
    run = para.add_run(text)
    if src_para.runs and src_para.runs[0]._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(src_para.runs[0]._r.rPr))
    return para


def build() -> bytes:
    src_bytes = SOURCE.read_bytes()
    if hashlib.sha256(src_bytes).hexdigest() != SOURCE_SHA256:
        die("時効側 送付案内.docx の SHA が想定と異なる（書式の正が変わった）")
    doc = Document(io.BytesIO(src_bytes))
    paras = doc.paragraphs
    if len(paras) < 20 or paras[_P_TITLE].text != "書　類　送　付　の　ご　案　内":
        die("時効側 送付案内.docx の段落構成が実測と異なる")

    # ── 冒頭ブロック（日付・宛先 3 行・署名）: プレースホルダを差し替え ──
    _set_single_run(paras[_P_DATE], "{{日付}}")
    _set_single_run(paras[_P_ADDR], "{{宛先郵便番号}}")
    addr_para = _clone_para(paras[_P_ADDR], "{{宛先住所}}", paras[_P_ADDR]._p)
    # 宛先名: 先頭の全角スペース run は書式として残し、2 番目の run をプレースホルダに
    name_runs = paras[_P_NAME].runs
    if len(name_runs) != 2:
        die("宛先名段落の run 構成が実測と異なる")
    name_runs[1].text = "{{宛先名}}"
    _set_single_run(paras[_P_SIGN], "{{事務所署名ブロック}}")
    # ── 表題 ──
    title = paras[_P_TITLE]
    _set_single_run(title, TITLE)
    # ── 本文: 元の 7 段落目以降と表を削除し、凍結本文を挿入 ──
    normal_src = copy.deepcopy(paras[_P_BODY_NORMAL])
    right_src = copy.deepcopy(paras[_P_KEIGU])
    center_src = copy.deepcopy(paras[_P_CENTER])
    body = doc.element.body
    for p in paras[_P_TITLE + 1:]:
        body.remove(p._p)
    for t in doc.tables:
        body.remove(t._tbl)
    # 表題の直後に空段落 1 つ（元書式と同じ）を置き、以降に本文
    from docx.text.paragraph import Paragraph
    anchor = _clone_para(Paragraph(normal_src._p, title._parent), "", title._p)
    prev = anchor
    for line in BODY_PARAGRAPHS:
        if line in RIGHT_ALIGNED:
            src = Paragraph(right_src._p, title._parent)
        elif line in CENTERED:
            src = Paragraph(center_src._p, title._parent)
        else:
            src = Paragraph(normal_src._p, title._parent)
        prev = _clone_para(src, line, prev._p)
        if line in RIGHT_ALIGNED:
            prev.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif line in CENTERED:
            prev.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # sectPr は body 末尾に残っている（段落削除は sectPr を触らない）
    buf = io.BytesIO()
    doc.save(buf)
    return deterministic_zip(buf.getvalue())


def verify(docx_bytes: bytes) -> None:
    doc = Document(io.BytesIO(docx_bytes))
    if doc.tables:
        die("表が残っている")
    texts = [p.text for p in doc.paragraphs]
    for key in PLACEHOLDERS:
        hits = [p for p in doc.paragraphs for r in p.runs if key in r.text]
        if not hits:
            die(f"プレースホルダ {key} が単一 run に存在しない")
    # 凍結本文: 表題以降の段落（先頭の空段落を除く）が BODY_PARAGRAPHS と一致
    i = texts.index(TITLE)
    body = texts[i + 2:]
    if tuple(body) != tuple(BODY_PARAGRAPHS):
        die("本文段落が凍結文と一致しない")
    if body_sha256(body) != BODY_SHA256:
        die("本文 sha256 が pin と一致しない")


def export_pdf(docx_path: Path, pdf_path: Path) -> bool:
    """Word COM で PDF 見本を出力（Word が無い環境は False）。"""
    import subprocess
    ps = (
        "$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
        f"$d=$w.Documents.Open('{docx_path}'); $d.SaveAs2('{pdf_path}', 17); "
        "$d.Close(); $w.Quit()"
    )
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                           capture_output=True, timeout=120)
        return r.returncode == 0 and pdf_path.is_file()
    except Exception:
        return False


def main() -> None:
    out = build()
    verify(out)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(out)
    print("wrote", OUT, "sha256", hashlib.sha256(out).hexdigest())
    if DESKTOP.is_dir():
        d = DESKTOP / "受理通知送付状_v1.docx"
        d.write_bytes(out)
        print("wrote", d)
        ok = export_pdf(d, DESKTOP / "受理通知送付状_v1.pdf")
        print("pdf", "ok" if ok else "skipped")


if __name__ == "__main__":
    main()
