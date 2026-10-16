"""送付状（家裁提出パッケージ）雛形の収載スクリプト（HOUKI-TEISHUTSU-1）

python-docx で白紙（A4・余白 25mm）から生成する。全文 MS Mincho（Normal スタイルの eastAsia フォント）・太字なし・
表・テキストボックス不使用の素の段落のみ。各 {{ }} は単一 run。生成物は決定的 zip（タイムスタンプ固定）で、SHA-256 を
hub/houki_teishutsu.TEMPLATE_SHA256 に pin する（代理人目録・上申書と同型）。

実行: python scripts/make_soufujo_template.py
"""

import hashlib
import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from hub.houki_teishutsu import FONT_JA, LINES, TEMPLATE_SHA256, verify_structure  # noqa: E402

OUT = REPO / "docx_templates" / "houki" / "送付状.docx"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_ALIGN = {"L": WD_ALIGN_PARAGRAPH.LEFT, "C": WD_ALIGN_PARAGRAPH.CENTER, "R": WD_ALIGN_PARAGRAPH.RIGHT}


def deterministic_zip(docx_bytes: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0
            dst.writestr(zi, src.read(info.filename))
    return out.getvalue()


def build() -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(25)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_JA
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_JA)
    for align, text in LINES:
        p = doc.add_paragraph()
        p.alignment = _ALIGN[align]
        if text:
            p.add_run(text)                      # 単一 run・書式指定なし（Normal を継承＝明朝・太字なし）
    buf = io.BytesIO()
    doc.save(buf)
    return deterministic_zip(buf.getvalue())


def main() -> None:
    out = build()
    verify_structure(Document(io.BytesIO(out)))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(out)
    # sink 方針: 値（パス・SHA）は出力しない。SHA は hub/houki_teishutsu.TEMPLATE_SHA256 と照合して確認する
    if hashlib.sha256(out).hexdigest() != TEMPLATE_SHA256:
        raise SystemExit("収載完了（SHA が hub/houki_teishutsu.TEMPLATE_SHA256 と不一致＝pin を更新して再実行）")
    print("wrote docx_templates/houki/送付状.docx (sha256 matches hub/houki_teishutsu.TEMPLATE_SHA256)")


if __name__ == "__main__":
    main()
