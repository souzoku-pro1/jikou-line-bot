# -*- coding: utf-8 -*-
"""時効援用完了報告書の雛形 docx を生成する — JIKOU-CERT-1（QUEUE-2026-10-20-HI I-1）

成果物: docx_templates/jikou/時効援用完了報告書.docx
規約（票）: 全文明朝体・太字なし・表やテキストボックス不使用の **素の段落** のみ。プレースホルダは各段落 1 つ（単一 run）。
複数行の値は使わない（contract_pdf.docx_to_pdf_bytes が段落単位で描画するため、明細・署名は 1 行 1 段落）。
zip の時刻を固定した決定的バイト列（同一定義→同一 SHA-256）。
使い方: python scripts/make_jikou_cert_template.py [<出力先>]
sink 規律: print は定数＋emit のみ。
"""
import hashlib
import io
import os
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from hub.redact import emit  # noqa: E402

_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FONT_JA = "ＭＳ 明朝"
OUT_DEFAULT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docx_templates", "jikou", "時効援用完了報告書.docx")

# 段落の並び（プレースホルダは 1 段落 1 つ・単一 run）
PARAGRAPHS = [
    ("{{表題}}", 14, "center"),
    ("", 10.5, "left"),
    ("{{作成日}}", 10.5, "right"),
    ("", 10.5, "left"),
    ("{{宛名}}", 10.5, "left"),
    ("", 10.5, "left"),
    ("{{本文}}", 10.5, "left"),
    ("", 10.5, "left"),
    ("{{明細1}}", 10.5, "left"),
    ("{{明細2}}", 10.5, "left"),
    ("{{明細3}}", 10.5, "left"),
    ("{{明細4}}", 10.5, "left"),
    ("", 10.5, "left"),
    ("", 10.5, "left"),
    ("{{署名1}}", 10.5, "right"),
    ("{{署名2}}", 10.5, "right"),
    ("{{署名3}}", 10.5, "right"),
    ("{{署名4}}", 10.5, "right"),
]
PLACEHOLDERS = tuple(p for p, _, _ in PARAGRAPHS if p)


def _set_font(run, size: float) -> None:
    run.font.name = FONT_JA
    run.font.size = Pt(size)
    run.font.bold = False
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_JA)


def build_docx() -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(25)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_JA
    normal.font.bold = False
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_JA)
    for text, size, align in PARAGRAPHS:
        para = doc.add_paragraph()
        para.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        run = para.add_run(text)
        _set_font(run, size)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def deterministic_zip(docx_bytes: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, src.read(info.filename))
    return buf.getvalue()


def main(argv) -> int:
    out = argv[1] if len(argv) > 1 else OUT_DEFAULT
    data = deterministic_zip(build_docx())
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "wb") as f:
        f.write(data)
    # sink 規律（test_sink_ast_policy）: print 引数は定数+emit() のみ（出力先 path は argv 由来のため出さない）
    print("written: docx_templates/jikou/時効援用完了報告書.docx (or argv[1])")
    print("TEMPLATE_SHA256 =", emit(hashlib.sha256(data).hexdigest(), "record_id", "log", "operator"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
