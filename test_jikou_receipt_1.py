"""JIKOU-RECEIPT-1（QUEUE-2026-10-20-HI I-2）: 領収書（xlsx）の採番・金額・発行・再発行・画面。

pin する契約:
  - 金額は料金規則から（手入力なし）: 1 社 44,000／2 社 88,000／3 社 110,000／追加 1 社（3 社目）22,000 の内訳一致・
    税込合計は chat_responder.fee_total_yen と一致・消費税 10% 切捨て（要大野確認）
  - 番号は 年+連番・並行 10 件で重複なし・再発行は同番号＋「再発行」表示（元の行は残る・請求済み社数には数えない）
  - 入金状況が「入金済み」でなければ発行不可・未請求 0 は発行不可・但し書きが承認台帳に無ければ発行不可
  - 雛形は決定的 zip で SHA pin・差し込み残りは拒否・数値セルは int + '#,##0'・登録番号は env があるときだけ
  - 画面: 確認なし POST は発行しない・ダウンロードは _gate の内側・ファイル名に氏名なし
  - 送信経路なし（hub/jikou_receipt は line_channel／push／notify を import しない）
"""

import asyncio
import datetime
import io
import os
import unittest
from unittest.mock import AsyncMock, patch

from test_jikou_reply_q1c import _DbMixin, _auth, _client, _run  # noqa: E402

import chat_responder  # noqa: E402
from hub import approved_texts as AT  # noqa: E402
from hub import db  # noqa: E402
from hub import jikou_doc_texts as T  # noqa: E402
from hub import jikou_followup as F  # noqa: E402
from hub import jikou_receipt as R  # noqa: E402
from hub import kintone  # noqa: E402
from hub import webapp_jikou_receipt_view as V  # noqa: E402

RID = "88"
OFFICE = {"OFFICE_NAME": "大野法律事務所", "OFFICE_ADDRESS": "川口市西青木2-1-45", "OFFICE_TEL": "048-299-2704"}


def _record(creditors=("A社",), paid="入金済み", name="山田 太郎") -> dict:
    cols = {"問い合わせ業者名": creditors[0] if len(creditors) > 0 else "", "対象債権者2": creditors[1] if len(creditors) > 1 else "",
            "対象債権者3": creditors[2] if len(creditors) > 2 else ""}
    base = {"顧客名": name, "入金状況": paid, "LINEユーザーID": "U" + "e" * 32, "電話番号": "090-0000-1111", **cols}
    return {"$id": {"value": RID}, **{k: {"value": v} for k, v in base.items()}}


def _approve():
    return patch.dict(AT.APPROVED_SHA256, {R.TEXT_KEY_NOTE: AT.sha256_text(T.text(R.TEXT_KEY_NOTE))})


def _sheet_values(xlsx: bytes) -> dict:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xlsx)).active
    return {c.coordinate: c for row in ws.iter_rows() for c in row if c.value not in (None, "")}


class TestCompute(unittest.TestCase):
    def test_breakdown_matches_fee_rule(self):
        self.assertTrue(R.consistent_with_fee_rule())
        for n, total in ((1, 44_000), (2, 88_000), (3, 110_000), (4, 132_000)):
            with self.subTest(n=n):
                b = R.compute(n)
                self.assertEqual(b.total, total)
                self.assertEqual(b.total, chat_responder.fee_total_yen(n))
        b = R.compute(3)
        self.assertEqual((b.qty1, b.unit1, b.amt1, b.qty2, b.unit2, b.amt2, b.subtotal, b.tax), (2, 40_000, 80_000, 1, 20_000, 20_000, 100_000, 10_000))

    def test_additional_creditor_is_difference_only(self):
        b = R.compute(3, already_billed=2)         # 2 社発行済み→3 社目を追加
        self.assertEqual((b.count_billed, b.qty1, b.qty2, b.subtotal, b.tax, b.total), (1, 0, 1, 20_000, 2_000, 22_000))
        b = R.compute(2, already_billed=1)         # 1 社→2 社目を追加
        self.assertEqual((b.count_billed, b.qty1, b.qty2, b.total), (1, 1, 0, 44_000))
        self.assertEqual(R.compute(2, already_billed=2).count_billed, 0)

    def test_number_format(self):
        self.assertEqual(R.format_number(2026, 7), "2026-0007")


class _Base(_DbMixin, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self._env = patch.dict(os.environ, OFFICE)
        self._env.start()
        self.addCleanup(self._env.stop)
        os.environ.pop(R.INVOICE_NO_ENV, None)

        async def _create():
            async with db.get_async_engine().begin() as c:
                await c.run_sync(F.metadata.create_all)
                await c.run_sync(R.metadata.create_all)
        _run(_create())
        self.record = _record()
        self._kp = patch.object(kintone, "get_record", AsyncMock(side_effect=lambda app, rid: self.record))
        self._kp.start()
        self.addCleanup(self._kp.stop)


class TestTemplate(unittest.TestCase):
    def test_pin_and_placeholders(self):
        self.assertTrue(R.template_ok())
        from openpyxl import load_workbook
        ws = load_workbook(R.TEMPLATE_PATH).active
        vals = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        for ph in ("{{領収書番号}}", "{{領収日}}", "{{宛名}}", "{{但し書き}}", "{{数量1}}", "{{単価1}}", "{{金額1}}", "{{数量2}}",
                   "{{単価2}}", "{{金額2}}", "{{小計}}", "{{消費税}}", "{{合計}}", "{{再発行}}", "{{登録番号}}", "{{債権者名}}"):
            self.assertTrue(any(ph in v for v in vals), ph)
        self.assertFalse(any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row))  # セル式なし

    def test_module_has_no_send_path(self):
        import inspect
        src = inspect.getsource(R)
        for bad in ("line_channel", "push_", "notify_"):
            self.assertNotIn(bad, src)


class TestIssue(_Base):
    def test_refusals(self):
        with self.assertRaises(R.IssueRefused) as ctx:
            _run(R.issue(RID, actor="human"))
        self.assertEqual(str(ctx.exception), "text_unapproved")
        with _approve():
            self.record = _record(paid="未入金")
            with self.assertRaises(R.IssueRefused) as ctx:
                _run(R.issue(RID, actor="human"))
            self.assertEqual(str(ctx.exception), "not_paid")
            self.record = _record(creditors=())
            with self.assertRaises(R.IssueRefused) as ctx:
                _run(R.issue(RID, actor="human"))
            self.assertEqual(str(ctx.exception), "no_creditors")
            self.record = _record(name="")
            with self.assertRaises(R.IssueRefused) as ctx:
                _run(R.issue(RID, actor="human"))
            self.assertEqual(str(ctx.exception), "name_missing")
        self.assertEqual(_run(R.list_receipts(RID)), [])

    def test_issue_amounts_sheet_and_difference(self):
        now = datetime.datetime(2026, 10, 21, 3, 0, tzinfo=datetime.timezone.utc)
        with _approve():
            self.record = _record(creditors=("A社", "B社"))
            out = _run(R.issue(RID, actor="human", now=now))
            self.assertEqual((out["number"], out["total"], out["count_billed"]), ("2026-0001", 88_000, 2))
            doc = _run(R.get_receipt(RID, out["receipt_id"]))
            cells = _sheet_values(doc["xlsx"])
            texts = {k: c.value for k, c in cells.items()}
            self.assertIn("No. 2026-0001", texts["E3"])
            self.assertEqual(texts["E4"], "令和8年10月21日")
            self.assertEqual(texts["B6"], "山田 太郎　様")
            self.assertEqual(cells["E16"].value, 88_000)
            self.assertEqual(cells["E16"].number_format, "#,##0")               # 半角・桁区切り（数値セル）
            self.assertEqual((cells["C12"].value, cells["D12"].value, cells["E12"].value), (2, 40_000, 80_000))
            self.assertEqual((cells["C13"].value, cells["E13"].value, cells["E15"].value), (0, 0, 8_000))
            self.assertIn("A社・B社", texts["B9"])                                 # 但し書きに対象債権者
            self.assertNotIn("B18", texts)                                      # 再発行表示なし
            self.assertNotIn("B23", texts)                                      # 登録番号欄なし（env 未設定）
            self.assertFalse(any("{{" in str(v) for v in texts.values()))
            self.assertNotIn("090-0000-1111", "".join(str(v) for v in texts.values()))
            # 追加依頼: 3 社目を追加 → 差額 22,000 のみ・新番号
            self.record = _record(creditors=("A社", "B社", "C社"))
            out2 = _run(R.issue(RID, actor="human", now=now))
            self.assertEqual((out2["number"], out2["total"], out2["count_billed"], out2["qty1"], out2["qty2"]), ("2026-0002", 22_000, 1, 0, 1))
            doc2 = _run(R.get_receipt(RID, out2["receipt_id"]))
            self.assertIn("C社", doc2["creditor_names"])
            self.assertNotIn("A社", doc2["creditor_names"])
            with self.assertRaises(R.IssueRefused) as ctx:                      # 未請求 0
                _run(R.issue(RID, actor="human", now=now))
            self.assertEqual(str(ctx.exception), "nothing_to_bill")

    def test_invoice_number_only_when_env_set(self):
        with _approve(), patch.dict(os.environ, {R.INVOICE_NO_ENV: "T1234567890123"}):
            out = _run(R.issue(RID, actor="human"))
        texts = {k: c.value for k, c in _sheet_values(_run(R.get_receipt(RID, out["receipt_id"]))["xlsx"]).items()}
        self.assertEqual(texts["B23"], "登録番号　T1234567890123")

    def test_reissue_same_number_with_label(self):
        with _approve():
            out = _run(R.issue(RID, actor="human"))
            re_ = _run(R.reissue(RID, out["receipt_id"], actor="human"))
        self.assertEqual(re_["number"], out["number"])
        self.assertTrue(re_["reissue"])
        rows = _run(R.list_receipts(RID))
        self.assertEqual([(r["number"], bool(r["reissue"])) for r in rows], [("2026-0001", False), ("2026-0001", True)])
        self.assertEqual(_run(R.billed_count(RID)), 1)                           # 再発行は請求済みに数えない
        texts = {k: c.value for k, c in _sheet_values(_run(R.get_receipt(RID, re_["receipt_id"]))["xlsx"]).items()}
        self.assertEqual(texts["B18"], "再発行")
        with self.assertRaises(R.IssueRefused):
            _run(R.reissue(RID, "0" * 32, actor="human"))

    def test_concurrent_allocation_has_no_duplicates(self):
        async def _many():
            return await asyncio.gather(*(R.allocate_number(2026) for _ in range(10)))
        results = _run(_many())
        numbers = [n for _, n in results]
        self.assertEqual(len(set(numbers)), 10)
        self.assertEqual(sorted(s for s, _ in results), list(range(1, 11)))


class TestView(_Base):
    def test_page_and_issue_flow(self):
        r = _client.get(f"{V.PAGE}/{RID}", headers=_auth())
        self.assertEqual(r.status_code, 200)
        self.assertIn("入金済み", r.text)
        self.assertIn("44,000", r.text)
        self.assertNotIn("山田", r.text)
        r = _client.post(f"{V.PAGE}/{RID}/issue", data={}, headers=_auth(), follow_redirects=False)
        self.assertIn("done=confirm_required", r.headers["location"])
        r = _client.post(f"{V.PAGE}/{RID}/issue", data={"confirm": "1"}, headers=_auth(), follow_redirects=False)
        self.assertIn("done=refused_text_unapproved", r.headers["location"])
        with _approve():
            r = _client.post(f"{V.PAGE}/{RID}/issue", data={"confirm": "1"}, headers=_auth(), follow_redirects=False)
        self.assertIn("done=issued", r.headers["location"])
        rows = _run(R.list_receipts(RID))
        self.assertEqual(len(rows), 1)
        dl = _client.get(f"{V.PAGE}/{RID}/doc/{rows[0]['receipt_id']}.xlsx", headers=_auth())
        self.assertEqual(dl.status_code, 200)
        self.assertTrue(dl.content.startswith(b"PK"))
        self.assertNotIn("山田", dl.headers["content-disposition"])
        with _approve():
            r = _client.post(f"{V.PAGE}/{RID}/reissue/{rows[0]['receipt_id']}", data={"confirm": "1"}, headers=_auth(), follow_redirects=False)
        self.assertIn("done=reissued", r.headers["location"])
        self.assertEqual(_client.get(f"{V.PAGE}/abc", headers=_auth()).status_code, 400)
        self.assertEqual(_client.get(f"{V.PAGE}/{RID}/doc/{'0' * 32}.xlsx", headers=_auth()).status_code, 404)

    def test_unauthenticated_is_rejected(self):
        r = _client.get(f"{V.PAGE}/{RID}", follow_redirects=False)
        self.assertIn(r.status_code, (303, 401, 404))


if __name__ == "__main__":
    unittest.main()
