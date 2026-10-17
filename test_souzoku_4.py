"""SOUZOKU-4（QUEUE-2026-10-17-H H-5）: 財産目録の構造化入力・分割案・遺産分割協議書（案）docx。

- 口座番号は下 4 桁以外を保存しない。生年月日の欄は無い。
- 分割案と法定相続分の差異表示。生成は人の明示操作のみ・App 添付なし（DB 保管→ダウンロード）。
- fixture: 相続人 3 名（配偶者＋子 2・確定済み）・不動産 1・預金 2 → docx が生成され、本文に口座番号全桁・生年月日が無い。
- 雛形は素の段落のみ（表・テキストボックスなし）・明朝・太字なし・透かし 1 行が先頭。条文は要承認表（pin）。
"""
import datetime
import io
import json
import os
import re
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import sqlalchemy as sa
from docx import Document

from test_souzoku_2 import GOOD, _run  # noqa: E402

REPO = Path(__file__).parent

import main  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from hub import db  # noqa: E402
from hub import send_ledger as sl  # noqa: E402
from hub import souzoku_assets as A  # noqa: E402
from hub import souzoku_heirs as H  # noqa: E402
from hub import souzoku_intake as S  # noqa: E402
from hub import souzoku_kyogisho as K  # noqa: E402
from hub import souzoku_texts as T  # noqa: E402
from hub import webapp_souzoku_assets_view as V  # noqa: E402
from hub.webapp_auth import MIN_ITERATIONS, hash_password, issue_session  # noqa: E402
from test_sink_ast_policy import scan_source  # noqa: E402

_client = TestClient(main.app)
FULL_ACCOUNT = "1234567"
FULL_ACCOUNT_2 = "9876543210"
BIRTH = "1960年4月1日"


def _auth():
    return {"Cookie": f"webapp_session={issue_session()}"}


class _Db(unittest.TestCase):
    def setUp(self):
        self._dir = tempfile.mkdtemp(prefix="souzoku4_")
        self._env = patch.dict(os.environ, {
            "DATABASE_URL": f"sqlite+aiosqlite:///{self._dir}/n.db", S.FLAG_ENV: "1",
            "WEBAPP_PASSWORD_HASH": hash_password("pw", iterations=MIN_ITERATIONS), "WEBAPP_SESSION_SECRET": "s" * 32})
        self._env.start()
        db.reset_for_tests()

        async def _create():
            async with db.get_async_engine().begin() as c:
                for md in (sl.metadata, S.metadata, H.metadata, A.metadata):
                    await c.run_sync(md.create_all)
        _run(_create())
        data, errors = S.validate(dict(GOOD, dispute="ない"))
        assert not errors, errors
        self.intake_id = _run(S.save(data))
        # 相続人 3 名（配偶者＋子 2）を確定済みに
        out = _run(H.save_draft(self.intake_id, {"spouse": "alive", "children": [{"status": "alive"}, {"status": "alive"}]}))
        assert _run(H.confirm(out["heirs_id"]))

    def tearDown(self):
        db.reset_for_tests()
        self._env.stop()
        shutil.rmtree(self._dir, ignore_errors=True)

    def url(self, tail="assets"):
        return f"/app/souzoku/{self.intake_id}/{tail}"

    def add(self, kind, value="", **fields):
        data = {"kind": kind, "value_yen": value}
        data.update({f"{kind}__{k}": v for k, v in fields.items()})
        return _client.post(self.url(), headers=_auth(), data=data, follow_redirects=False)

    def assets(self):
        return _run(A.list_assets(self.intake_id, include_void=True))

    def fixture(self):
        self.add("real_estate", "30000000", location="川口市西青木二丁目", lot_number="1番45", house_number="45", share="1/1")
        self.add("deposit", "5000000", bank="みずほ銀行", branch="川口支店", account_type="普通", account_number=FULL_ACCOUNT)
        self.add("deposit", "", bank="ゆうちょ銀行", branch="〇一八", account_type="通常貯金", account_number=FULL_ACCOUNT_2)
        return self.assets()


class TestAssets(_Db):
    def test_flag_gate_and_404(self):
        with patch.dict(os.environ, {S.FLAG_ENV: ""}):
            self.assertEqual(_client.get(self.url(), headers=_auth()).status_code, 404)
        self.assertEqual(_client.get(self.url(), follow_redirects=False).status_code, 303)
        for r in V.router.routes:
            self.assertTrue(getattr(r.endpoint, "__webapp_gate__", False), r.path)
        self.assertEqual(_client.get("/app/souzoku/" + "f" * 32 + "/assets", headers=_auth()).status_code, 404)

    def test_add_stores_only_last4_of_account_and_no_birthdate_field(self):
        rows = self.fixture()
        self.assertEqual([a["kind"] for a in rows], ["real_estate", "deposit", "deposit"])
        dep = rows[1]["detail"]
        self.assertEqual(dep["account_last4"], "4567")
        self.assertNotIn("account_number", dep)
        for a in rows:
            blob = json.dumps(a["detail"], ensure_ascii=False)
            self.assertNotIn(FULL_ACCOUNT, blob)
            self.assertNotIn(FULL_ACCOUNT_2, blob)
            self.assertNotIn("生年月日", blob)
        self.assertEqual(rows[2]["value_yen"], None)                       # 評価未確定
        self.assertEqual(A.account_last4("12-34567"), "4567")
        self.assertEqual(A.account_last4(""), "")
        self.assertEqual({k for f in A.FIELDS.values() for k, _ in f} & {"birth", "birthday", "account_number"}, set())
        page = _client.get(self.url(), headers=_auth()).text
        self.assertIn("下4桁 4567", page)
        self.assertNotIn(FULL_ACCOUNT, page)
        self.assertIn("弁護士確認前", page)
        # 不正入力・空
        self.assertIn("done=bad_input", self.add("deposit", "abc", bank="x").headers["location"])
        self.assertIn("done=bad_input", self.add("other").headers["location"])
        self.assertEqual(len(self.assets()), 3)

    def test_allocate_void_and_diff(self):
        rows = self.fixture()
        re_id, d1, d2 = [a["asset_id"] for a in rows]
        r = _client.post(f"/app/souzoku/{self.intake_id}/assets/{re_id}/allocate", headers=_auth(),
                         data={"acquirer": "配偶者"}, follow_redirects=False)
        self.assertEqual(r.status_code, 303, r.text[:200])
        self.assertIn("done=allocated", r.headers["location"])
        _client.post(f"/app/souzoku/{self.intake_id}/assets/{d1}/allocate", headers=_auth(),
                     data={"acquirer": "子1", "comp_to": "子2", "comp_yen": "1000000", "comp_due": "2026-12-31"})
        # 不正: 判定に無いラベル・自分への代償金
        r = _client.post(f"/app/souzoku/{self.intake_id}/assets/{d2}/allocate", headers=_auth(), data={"acquirer": "他人"}, follow_redirects=False)
        self.assertIn("done=bad_input", r.headers["location"])
        r = _client.post(f"/app/souzoku/{self.intake_id}/assets/{d1}/allocate", headers=_auth(),
                         data={"acquirer": "子1", "comp_to": "子1", "comp_yen": "1"}, follow_redirects=False)
        self.assertIn("done=bad_input", r.headers["location"])
        d = A.diff_table(_run(H.latest(self.intake_id))["result_detail"], _run(A.list_assets(self.intake_id)))
        self.assertEqual((d["total_yen"], d["unvalued"], d["unallocated"]), (35_000_000, 1, 1))
        by = {r["label"]: r for r in d["rows"]}
        self.assertEqual((by["配偶者"]["share"], by["配偶者"]["expected_yen"], by["配偶者"]["acquired_yen"]), ("1/2", 17_500_000, 30_000_000))
        self.assertEqual((by["子1"]["acquired_yen"], by["子2"]["acquired_yen"]), (4_000_000, 1_000_000))
        self.assertEqual(by["子2"]["diff_yen"], 1_000_000 - 8_750_000)
        page = _client.get(self.url(), headers=_auth()).text
        self.assertIn("法定相続分との差異", page)
        self.assertIn("17,500,000 円", page)
        # 無効化（削除しない）
        r = _client.post(f"/app/souzoku/{self.intake_id}/assets/{d2}/void", headers=_auth(), follow_redirects=False)
        self.assertIn("done=voided", r.headers["location"])
        self.assertEqual([a["status"] for a in self.assets()], ["active", "active", "void"])
        self.assertEqual(len(_run(A.list_assets(self.intake_id))), 2)
        r = _client.post(f"/app/souzoku/{self.intake_id}/assets/{d2}/void", headers=_auth(), follow_redirects=False)
        self.assertIn("done=not_found", r.headers["location"])


class TestKyogisho(_Db):
    def _allocate_all(self, rows):
        re_id, d1, d2 = [a["asset_id"] for a in rows]
        _client.post(f"/app/souzoku/{self.intake_id}/assets/{re_id}/allocate", headers=_auth(), data={"acquirer": "配偶者"})
        _client.post(f"/app/souzoku/{self.intake_id}/assets/{d1}/allocate", headers=_auth(),
                     data={"acquirer": "子1", "comp_to": "子2", "comp_yen": "1000000", "comp_due": "2026-12-31"})
        _client.post(f"/app/souzoku/{self.intake_id}/assets/{d2}/allocate", headers=_auth(), data={"acquirer": "子2"})

    def test_generate_requires_confirm_and_produces_docx_without_full_account_or_birthdate(self):
        rows = self.fixture()
        self._allocate_all(rows)
        r = _client.post(self.url("kyogisho"), headers=_auth(), data={}, follow_redirects=False)
        self.assertIn("done=confirm_required", r.headers["location"])
        self.assertEqual(_run(K.list_docs(self.intake_id)), [])
        r = _client.post(self.url("kyogisho"), headers=_auth(), data={"confirm": "1"}, follow_redirects=False)
        self.assertIn("done=generated", r.headers["location"])
        docs = _run(K.list_docs(self.intake_id))
        self.assertEqual([d["version"] for d in docs], [1])
        dl = _client.get(f"/app/souzoku/{self.intake_id}/kyogisho/{docs[0]['doc_id']}.docx", headers=_auth())
        self.assertEqual(dl.status_code, 200)
        self.assertIn("wordprocessingml", dl.headers["content-type"])
        self.assertEqual(dl.headers.get("cache-control"), "no-store, private")
        doc = Document(io.BytesIO(dl.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        # 本文: 透かしが先頭・被相続人・相続人 3 名・財産 3・代償金・口座は下 4 桁のみ・生年月日なし
        self.assertTrue(text.startswith(T.text("K-0")))
        self.assertIn("山田 太郎", text)                                  # 被相続人（受付の値）
        for lb in ("配偶者", "子1", "子2"):
            self.assertIn(f"相続人（{lb}）", text)
        self.assertIn("下4桁 4567", text)
        self.assertNotIn(FULL_ACCOUNT, text)
        self.assertNotIn(FULL_ACCOUNT_2, text)
        self.assertIsNone(re.search(r"\d{5,}", text))
        self.assertNotIn(BIRTH, text)
        self.assertNotIn("生年月日", text)
        self.assertIn("金1,000,000円", text)
        self.assertIn("氏名 ＿＿", text)                                   # 氏名・住所は空欄線（弁護士が記入）
        self.assertIn(T.text("K-11"), text)                              # 評価未確定の注記
        self.assertEqual(len(doc.tables), 0)                             # 表なし
        self.assertTrue(all(not r.font.bold for p in doc.paragraphs for r in p.runs))
        # 2 回目は版 2（上書きしない）・同じ入力なら同じ SHA（決定的）
        _client.post(self.url("kyogisho"), headers=_auth(), data={"confirm": "1"})
        docs = _run(K.list_docs(self.intake_id))
        self.assertEqual([d["version"] for d in docs], [2, 1])
        self.assertEqual(docs[0]["sha256"], docs[1]["sha256"])
        # 送信・通知なし
        async def _ops():
            async with db.session_scope() as s:
                return (await s.execute(sa.select(sa.func.count()).select_from(sl.send_operation))).scalar()
        self.assertEqual(_run(_ops()), 0)
        page = _client.get(self.url(), headers=_auth()).text
        self.assertIn("ダウンロード", page)
        self.assertEqual(_client.get(f"/app/souzoku/{self.intake_id}/kyogisho/{'0' * 32}.docx", headers=_auth()).status_code, 404)

    def test_generation_refused_without_heirs_or_with_bad_pin(self):
        other_data, _ = S.validate(dict(GOOD, dispute="ない"))
        other = _run(S.save(other_data))
        r = _client.post(f"/app/souzoku/{other}/kyogisho", headers=_auth(), data={"confirm": "1"}, follow_redirects=False)
        self.assertIn("done=no_heirs", r.headers["location"])
        self.fixture()
        with patch.object(T, "TABLE_SHA256_PIN", "0" * 64):
            r = _client.post(self.url("kyogisho"), headers=_auth(), data={"confirm": "1"}, follow_redirects=False)
        self.assertIn("done=refused", r.headers["location"])
        with patch.object(K, "TEMPLATE_SHA256", "0" * 64):
            r = _client.post(self.url("kyogisho"), headers=_auth(), data={"confirm": "1"}, follow_redirects=False)
        self.assertIn("done=refused", r.headers["location"])
        self.assertEqual(_run(K.list_docs(self.intake_id)), [])
        # 未割付の注記（割付なしでも生成はできる＝案）
        r = _client.post(self.url("kyogisho"), headers=_auth(), data={"confirm": "1"}, follow_redirects=False)
        self.assertIn("done=generated", r.headers["location"])
        d = _run(K.get_doc(self.intake_id, _run(K.list_docs(self.intake_id))[0]["doc_id"]))
        text = "\n".join(p.text for p in Document(io.BytesIO(bytes(d["content"]))).paragraphs)
        self.assertIn("取得者が未定", text)
        self.assertIn("取得条項なし", text)


class TestPins(unittest.TestCase):
    def test_template_texts_migration_wiring_sink(self):
        import hashlib
        p = REPO / "docx_templates" / "souzoku" / "遺産分割協議書.docx"
        self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(), K.TEMPLATE_SHA256)
        doc = Document(str(p))
        self.assertEqual(len(doc.tables), 0)
        texts = [q.text for q in doc.paragraphs]
        self.assertEqual(texts[0], "{{透かし}}")
        for ph in K.PLACEHOLDERS:
            self.assertEqual(sum(ph in t for t in texts), 1, ph)
        for q in doc.paragraphs:
            for r in q.runs:
                self.assertFalse(bool(r.font.bold))
                self.assertEqual(r.font.name, "ＭＳ 明朝")
        self.assertNotIn("txbx", p.read_bytes().decode("latin-1"))             # テキストボックスなし
        self.assertTrue(T.verify_pin())
        self.assertTrue(all(e["status"] == T.STATUS_NEEDS_APPROVAL for e in T.TEXTS.values()))   # 全条文 要承認
        self.assertFalse(T.all_approved())
        m = REPO / "alembic" / "versions" / "20261017_d2e4f6a8b0c3_souzoku_assets.py"
        src = m.read_text(encoding="utf-8")
        self.assertIn("revision: str = 'd2e4f6a8b0c3'", src)
        self.assertIn("down_revision: Union[str, Sequence[str], None] = 'c1d3e5f7a9b2'", src)
        self.assertIn("souzoku_assets_downgrade_refused_data_present", src)
        self.assertIn("souzoku_assets_metadata", (REPO / "alembic" / "env.py").read_text(encoding="utf-8"))
        msrc = (REPO / "main.py").read_text(encoding="utf-8").splitlines()
        self.assertEqual(sum("include_before_catch_all(app, souzoku_assets_router)" in l for l in msrc), 1)
        for name in ("hub/souzoku_assets.py", "hub/souzoku_kyogisho.py", "hub/souzoku_texts.py", "hub/webapp_souzoku_assets_view.py"):
            # scripts/ は sink 方針の走査対象外（雛形生成スクリプトの print はバイト数と SHA のみ）
            self.assertEqual(scan_source((REPO / name).read_text(encoding="utf-8"), name), [], name)
        allow = (REPO / "redaction_sink_allowlist.json").read_text(encoding="utf-8")
        self.assertNotIn("souzoku_assets", allow)
        self.assertNotIn("souzoku_kyogisho", allow)
        # 外部書込・送信の口が無い
        for name in ("hub/souzoku_assets.py", "hub/souzoku_kyogisho.py", "hub/webapp_souzoku_assets_view.py"):
            body = (REPO / name).read_text(encoding="utf-8")
            for bad in ("import kintone", "hub.kintone", "kintone.", "line_channel", "push_text", "notify_admin_line", "upload_file"):
                self.assertNotIn(bad, body, (name, bad))


if __name__ == "__main__":
    unittest.main()
