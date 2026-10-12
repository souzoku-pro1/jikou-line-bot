"""HOUKI-JIJOU-1（QUEUE-2026-10-13-D D-5）: 事情説明書（上申書）の LLM 下書き・素材一致検査・二重承認・docx 生成

固定する仕様:
- 素材一致検査（check_against_source）: 下書きの日付（西暦・和暦）・年月・金額・固有名詞（カタカナ語・会社名）が
  App 40 の素材欄（SOURCE_FIELDS）に無ければ違反（固定語彙 "節:種別:値"）。素材にある値は通る。
- 承認なしで生成されない: GET は下書き表示のみ（kintone 書込ゼロ・upload なし）。POST は ticket・明示確認・revision
  一致・違反なし・FILE 欄あり のときだけ生成・添付（$revision CAS）。添付済みは作用 0。
- SHA: 雛形 docx_templates/houki/上申書_熟慮期間.docx の SHA-256 pin・改変で生成拒否。構造（表題・見出し・結語・
  プレースホルダ各 1）を検査。生成物にプレースホルダ残存なし。
- PII: 画面に LINE userId・電話・住所を出さない。ログ・通知は No.・件数・固定語彙。新 module の sink 違反ゼロ。
"""

import hashlib
import io
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document

for _k, _v in {
    "KINTONE_SUBDOMAIN": "testsub", "LINE_CHANNEL_SECRET": "dummy_secret",
    "LINE_CHANNEL_ACCESS_TOKEN": "dummy_token", "ANTHROPIC_API_KEY": "dummy_key",
    "KINTONE_APP_ID": "21", "KINTONE_API_TOKEN": "dummy",
    "SOUZOKU_KINTONE_APP_ID": "26", "SOUZOKU_KINTONE_API_TOKEN": "dummy",
    "CLOUDSIGN_CLIENT_ID": "c", "CLOUDSIGN_WEBHOOK_SECRET": "cs",
    "KINTONE_WEBHOOK_TOKEN": "kintone-token", "DOCUMENT_WEBHOOK_SECRET": "d",
    "APP_APPROVAL": "29", "TOKEN_APPROVAL": "d", "HEALTHCHECK_DISABLED": "1",
    "STRIPE_WEBHOOK_SECRET": "w", "GOOGLE_VISION_API_KEY": "dummy_vision",
    "APP_HOUKI": "40", "TOKEN_HOUKI": "t40",
}.items():
    os.environ.setdefault(_k, _v)

from fastapi.testclient import TestClient  # noqa: E402

import main  # noqa: E402
from hub import houki_jijou as J  # noqa: E402
from hub import kintone  # noqa: E402
from hub import webapp_houki_jijou_view as V  # noqa: E402
from hub.webapp_auth import MIN_ITERATIONS, hash_password, issue_session  # noqa: E402
from test_sink_ast_policy import scan_source  # noqa: E402

REPO = Path(__file__).parent
_client = TestClient(main.app)
_ENV = {"WEBAPP_PASSWORD_HASH": hash_password("pw", iterations=MIN_ITERATIONS),
        "WEBAPP_SESSION_SECRET": "s" * 32, "OFFICE_ATTORNEY": "大野 太郎"}
USER = "U" + "a" * 32


def _rec(**fields):
    out = {k: {"value": v} for k, v in fields.items()}
    out.setdefault("$id", {"value": "40"})
    out.setdefault("$revision", {"value": "3"})
    return out


RECORD = _rec(顧客名="山田 花子", 被相続人氏名="山田 太郎", 続柄="子", 死亡日_申告="2026-03-10",
              死亡を知った日_申告="2026-08-01", 知った経緯="債権者のアビリオ債権回収から 2026-08-01 に通知書が届いて知った。",
              財産_負債="アビリオ債権回収 約 1,200,000 円", 財産処分有無="なし", 訴訟督促有無="あり",
              起算日_確定="2026-08-01", 起算点メモ="死亡を知ったのは通知書到達日。被相続人とは 20 年間音信不通。",
              管轄家庭裁判所="さいたま家庭裁判所", 電話番号="090-1111-2222", LINEユーザーID=USER, 住所="埼玉県○○市")
GOOD = {"kankei": "申述人は被相続人の子です。被相続人とは 20 年間音信不通でした。",
        "keii": "令和8年8月1日にアビリオ債権回収からの通知書が届き、被相続人の死亡（2026年3月10日）を知りました。",
        "riyu": "被相続人の債務を相続することを望まないためです。",
        "saimu": "通知書により約 1,200,000 円の債務があると認識しています。",
        "shobun": "遺産の処分は行っていません。"}


def _auth():
    return {"Cookie": f"webapp_session={issue_session()}"}


class TestGuard(unittest.TestCase):
    def test_values_in_source_pass(self):
        self.assertEqual(J.check_against_source(GOOD, RECORD), [])

    def test_unknown_date_amount_noun_are_violations(self):
        bad = dict(GOOD, keii="2026年9月15日に知りました。", saimu="債務は 300万円 です。プロミス株式会社からも請求があります。")
        v = J.check_against_source(bad, RECORD)
        self.assertIn("keii:date:2026-09-15", v)
        self.assertIn("saimu:amount:3000000", v)
        self.assertTrue(any(x.startswith("saimu:noun:") and "プロミス" in x for x in v))
        self.assertTrue(any(x == "keii:month:2026-09" for x in v))

    def test_wareki_and_month_normalization(self):
        self.assertEqual(J._dates("令和8年8月1日・平成30年2月3日・2026/03/10"), {"2026-08-01", "2018-02-03", "2026-03-10"})
        self.assertEqual(J._months("2026年3月に死亡"), {"2026-03"})
        self.assertEqual(J._amounts("約 1,200,000 円と 30万円"), {1200000, 300000})
        self.assertEqual(J.check_against_source({"keii": "2026年3月に死亡しました。"}, RECORD), [])
        self.assertEqual(J.check_against_source({k: "" for k in J.SECTION_KEYS}, RECORD), [])

    def test_source_fields_closed_and_pii_fields_excluded(self):
        self.assertNotIn("電話番号", J.SOURCE_FIELDS)
        self.assertNotIn("LINEユーザーID", J.SOURCE_FIELDS)
        self.assertNotIn("住所", J.SOURCE_FIELDS)
        self.assertNotIn(USER, J.source_text(RECORD))
        self.assertIsNone(J.parse_draft({"kankei": "x"}))
        self.assertEqual(J.parse_draft({k: " a " for k in J.SECTION_KEYS})["kankei"], "a")
        with self.assertRaises(J.JijouGuardError):
            J.generate(RECORD, dict(GOOD, keii="2030年1月1日"))


class TestTemplate(unittest.TestCase):
    def test_template_pin_structure_and_render(self):
        data = Path(J.TEMPLATE_PATH).read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(), J.TEMPLATE_SHA256)
        doc = Document(io.BytesIO(data))
        J.verify_structure(doc)
        texts = "\n".join(p.text for p in J._all_paragraphs(doc))
        for ph in J.PLACEHOLDERS:
            self.assertEqual(texts.count(ph), 1, ph)
        with patch.dict(os.environ, {"OFFICE_ATTORNEY": "大野 太郎"}):
            out = J.generate(RECORD, GOOD, today=__import__("datetime").date(2026, 10, 13))
        body = "\n".join(p.text for p in J._all_paragraphs(Document(io.BytesIO(out))))
        self.assertNotIn("{{", body)
        self.assertIn("被相続人　山田 太郎　相続放棄申述事件", body)
        self.assertIn("さいたま家庭裁判所", body)
        self.assertIn(GOOD["kankei"], body)
        self.assertIn("令和8年10月13日", body)
        self.assertIn("大野 太郎", body)
        self.assertIn(J.CLOSING, body)
        # 空の節は【弁護士確認】で埋める（プレースホルダは残さない）
        fill = J.build_fill(RECORD, {"kankei": "x"})
        self.assertEqual(fill["{{経緯}}"], J.LAWYER_CHECK)

    def test_tampered_template_is_rejected(self):
        d = tempfile.mkdtemp(prefix="jijou_")
        try:
            p = Path(d) / "t.docx"
            data = bytearray(Path(J.TEMPLATE_PATH).read_bytes())
            data[-1] ^= 0x01
            p.write_bytes(bytes(data))
            with self.assertRaises(J.JijouIntegrityError):
                J.render(J.build_fill(RECORD, GOOD), template_path=p)
            with patch.object(J, "TEMPLATE_SHA256", "0" * 64):
                with self.assertRaises(J.JijouIntegrityError):
                    J.generate(RECORD, GOOD)
        finally:
            shutil.rmtree(d, ignore_errors=True)
        with self.assertRaises(J.JijouIntegrityError):
            J.verify_structure(Document(), filled=None)


class TestView(unittest.TestCase):
    def setUp(self):
        self._env = patch.dict(os.environ, _ENV)
        self._env.start()
        V.clear_caches()
        self.record = {k: dict(v) for k, v in RECORD.items()}
        self.get_record = AsyncMock(side_effect=lambda app, rid: self.record)
        self.upload = AsyncMock(return_value="fk1")
        self.update = AsyncMock()
        self.draft = AsyncMock(return_value=dict(GOOD))
        self.fields = AsyncMock(return_value={J.FIELD_ATTACHMENT: {"type": "FILE"}})
        self._p = [patch.object(kintone, "get_record", self.get_record), patch.object(kintone, "upload_file", self.upload),
                   patch.object(kintone, "update_record", self.update), patch.object(J, "draft_sections", self.draft),
                   patch.object(kintone, "get_form_fields", self.fields)]
        for p in self._p:
            p.start()

    def tearDown(self):
        for p in self._p:
            p.stop()
        V.clear_caches()
        self._env.stop()

    def _form(self):
        r = _client.get(f"{V.PAGE}/40", headers=_auth())
        self.assertEqual(r.status_code, 200)
        op = r.text.split('name="op_id" value="')[1].split('"')[0]
        tk = r.text.split('name="ticket" value="')[1].split('"')[0]
        return r, {"op_id": op, "ticket": tk, **GOOD}

    def test_gate_and_get_shows_draft_without_pii_and_writes_nothing(self):
        self.assertEqual(_client.get(f"{V.PAGE}/40", follow_redirects=False).status_code, 303)
        r, _ = self._form()
        self.assertIn(GOOD["kankei"], r.text)
        self.assertIn("承認して生成", r.text)
        self.assertNotIn(USER, r.text)
        self.assertNotIn("090-1111-2222", r.text)
        self.assertNotIn("埼玉県○○市", r.text)
        self.upload.assert_not_called()
        self.update.assert_not_called()
        self.assertEqual(self.draft.await_count, 1)
        _client.get(f"{V.PAGE}/40", headers=_auth())
        self.assertEqual(self.draft.await_count, 1)                 # 30 分キャッシュ
        self.assertEqual(_client.get(f"{V.PAGE}/abc", headers=_auth()).status_code, 400)

    def test_generate_requires_confirm_ticket_and_no_violations(self):
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=f, headers=_auth(), follow_redirects=False)
        self.assertIn("confirm_required", r.headers["location"])
        self.upload.assert_not_called()
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1", ticket="bad"), headers=_auth(),
                         follow_redirects=False)
        self.assertEqual(r.status_code, 400)
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1", keii="2030年1月1日に知った。"),
                         headers=_auth(), follow_redirects=False)
        self.assertEqual(r.status_code, 200)
        self.assertIn("素材一致検査の違反", r.text)
        self.assertIn("2030年1月1日に知った。", r.text)             # 修正欄は保持
        self.upload.assert_not_called()
        self.update.assert_not_called()

    def test_generate_attaches_once_with_cas(self):
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=generated", r.headers["location"])
        self.assertEqual(self.upload.await_count, 1)
        self.assertTrue(self.upload.await_args.args[1].endswith(".docx"))
        self.assertEqual(self.update.await_args.args[2], {J.FIELD_ATTACHMENT: [{"fileKey": "fk1"}]})
        self.assertEqual(self.update.await_args.kwargs["revision"], "3")
        # 添付済み → 作用 0
        self.record[J.FIELD_ATTACHMENT] = {"value": [{"fileKey": "fk1"}]}
        _, f2 = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f2, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=already", r.headers["location"])
        self.assertEqual(self.upload.await_count, 1)

    def test_stale_field_missing_and_attach_failure(self):
        _, f = self._form()
        self.record["$revision"] = {"value": "4"}
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=stale", r.headers["location"])
        self.upload.assert_not_called()
        V.clear_caches()
        self.fields.return_value = {}
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=field_missing", r.headers["location"])
        self.upload.assert_not_called()
        V.clear_caches()
        self.fields.return_value = {J.FIELD_ATTACHMENT: {"type": "FILE"}}
        self.update.side_effect = kintone.KintoneConflict("409")
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=stale", r.headers["location"])
        self.update.side_effect = RuntimeError("boom")
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=attach_failed", r.headers["location"])

    def test_draft_failure_shows_empty_sections(self):
        self.draft.side_effect = RuntimeError("api")
        r, _ = self._form()
        self.assertIn("下書きを生成できませんでした", r.text)
        self.assertIn("承認して生成", r.text)


class TestPins(unittest.TestCase):
    def test_sink_wiring_and_template_script(self):
        allow = (REPO / "redaction_sink_allowlist.json").read_text(encoding="utf-8")
        for name in ("hub/houki_jijou.py", "hub/webapp_houki_jijou_view.py", "scripts/make_jijou_template.py"):
            self.assertEqual(scan_source((REPO / name).read_text(encoding="utf-8"), name), [], name)
            self.assertNotIn(name.split("/")[-1], allow)
        msrc = (REPO / "main.py").read_text(encoding="utf-8").split("\n")
        self.assertIn("include_before_catch_all(app, houki_jijou_router)", "\n".join(msrc[-6:]))
        self.assertEqual(J.SECTION_KEYS, ("kankei", "keii", "riyu", "saimu", "shobun"))
        self.assertEqual(J.FIELD_ATTACHMENT, "事情説明書")


if __name__ == "__main__":
    unittest.main()
