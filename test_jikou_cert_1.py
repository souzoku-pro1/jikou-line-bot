"""JIKOU-CERT-1（QUEUE-2026-10-20-HI I-1）: 時効援用完了報告書（docx＋PDF）の生成・保管・画面。

pin する契約:
  - 成立（established）以外・発送日／回答日の欠落・債権者名なし・宛名なし は生成不可（固定理由）
  - 文面（hub/jikou_doc_texts）は承認台帳（hub/approved_texts）に本文 SHA が載るまで生成不可（text_unapproved）
  - 本文に LINE ユーザー ID・相談者の電話番号・メールアドレスが入らない（App 21 の値で照合・入れば pii_in_body）
  - 再生成で版が増える（1, 2, …）・上書きしない・docx と PDF の両方を保管・PDF 失敗は保存しない
  - 雛形は決定的 zip で SHA pin・素の段落のみ（表なし）
  - 画面: 確認なし POST は生成しない・成立の枠だけ生成ボタン・ダウンロードは _gate の内側・ファイル名に氏名なし
  - 自動生成なし・相談者への送信経路なし（hub/jikou_cert は line_channel／push を import しない）
"""

import asyncio
import datetime
import os
import re
import unittest
from unittest.mock import AsyncMock, patch

import sqlalchemy as sa

from test_jikou_reply_q1c import _DbMixin, _auth, _client, _run  # noqa: E402

from hub import approved_texts as AT  # noqa: E402
from hub import db  # noqa: E402
from hub import jikou_cert as C  # noqa: E402
from hub import jikou_doc_texts as T  # noqa: E402
from hub import jikou_followup as F  # noqa: E402
from hub import kintone  # noqa: E402
from hub import webapp_jikou_cert_view as V  # noqa: E402

RID = "77"
LINE_ID = "U" + "f" * 32
PHONE = "090-1234-5678"
MAIL = "taro@example.com"
OFFICE = {"OFFICE_NAME": "大野法律事務所", "OFFICE_ADDRESS": "川口市西青木2-1-45", "OFFICE_TEL": "048-299-2704",
          "OFFICE_ATTORNEY": "大野 太郎", "OFFICE_BAR": "埼玉弁護士会", "OFFICE_REG": "12345"}


def _record(**over) -> dict:
    base = {"顧客名": "山田 太郎", "LINEユーザーID": LINE_ID, "電話番号": PHONE, "メールアドレス": MAIL,
            "問い合わせ業者名": "A社", "対象債権者2": "B社", "対象債権者3": ""}
    base.update(over)
    return {"$id": {"value": RID}, **{k: {"value": v} for k, v in base.items()}}


def _approve_all():
    return patch.dict(AT.APPROVED_SHA256, {k: AT.sha256_text(T.text(k)) for k in T.TEXT_IDS})


class _Base(_DbMixin, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self._env = patch.dict(os.environ, OFFICE)
        self._env.start()
        self.addCleanup(self._env.stop)

        async def _create():
            async with db.get_async_engine().begin() as c:
                await c.run_sync(F.metadata.create_all)
                await c.run_sync(C.metadata.create_all)
        _run(_create())
        self.record = _record()
        self._kp = patch.object(kintone, "get_record", AsyncMock(side_effect=lambda app, rid: self.record))
        self._kp.start()
        self.addCleanup(self._kp.stop)
        _run(F.ensure_rows(RID, [(1, "A社"), (2, "B社")]))
        _run(F.set_shipped(RID, datetime.date(2026, 9, 1), source="human", actor="human"))
        _run(F.record_result(RID, 1, result=F.RESULT_ESTABLISHED, answered_on=datetime.date(2026, 10, 1), summary="時効完成を認める旨"))
        _run(F.record_result(RID, 2, result=F.RESULT_NOT_ESTABLISHED, answered_on=datetime.date(2026, 10, 2), summary=None))

    def row(self, slot):
        return {int(r["slot"]): r for r in _run(F.rows_for_case(RID))}[slot]


class TestTemplateAndTexts(unittest.TestCase):
    def test_template_pin_and_plain_paragraphs(self):
        from docx import Document
        self.assertTrue(C.template_ok())
        d = Document(C.TEMPLATE_PATH)
        self.assertEqual(d.tables, [])
        texts = [p.text for p in d.paragraphs if p.text]
        self.assertEqual(len(texts), len(set(texts)))                        # 1 段落 1 プレースホルダ
        for p in d.paragraphs:
            for r in p.runs:
                self.assertFalse(r.font.bold)
        self.assertTrue(T.verify_pin())
        self.assertTrue(all(e["status"] == T.STATUS_NEEDS_APPROVAL for e in T.TEXTS.values()))   # 全文言が要承認（仮置き）

    def test_module_has_no_send_path(self):
        import inspect
        src = inspect.getsource(C)
        self.assertNotIn("line_channel", src)
        self.assertNotIn("push_", src)
        self.assertNotIn("notify_", src)


class TestCompose(_Base):
    def test_refuses_without_approval_ledger(self):
        with self.assertRaises(C.GenerationRefused) as ctx:
            C.compose(self.record, self.row(1), today=datetime.date(2026, 10, 21))
        self.assertEqual(str(ctx.exception), "text_unapproved")

    def test_refuses_not_established_and_missing(self):
        with _approve_all():
            with self.assertRaises(C.GenerationRefused) as ctx:
                C.compose(self.record, self.row(2), today=datetime.date(2026, 10, 21))
            self.assertEqual(str(ctx.exception), "not_established")
            r = dict(self.row(1)); r["answered_on"] = None
            with self.assertRaises(C.GenerationRefused) as ctx:
                C.compose(self.record, r)
            self.assertEqual(str(ctx.exception), "dates_missing")
            r = dict(self.row(1)); r["creditor_name"] = ""
            with self.assertRaises(C.GenerationRefused) as ctx:
                C.compose(self.record, r)
            self.assertEqual(str(ctx.exception), "creditor_missing")
            with self.assertRaises(C.GenerationRefused) as ctx:
                C.compose(_record(顧客名=""), self.row(1))
            self.assertEqual(str(ctx.exception), "name_missing")

    def test_compose_values_and_wareki(self):
        with _approve_all():
            data = C.compose(self.record, self.row(1), today=datetime.date(2026, 10, 21))
        self.assertEqual(data["{{表題}}"], "時効援用完了報告書")
        self.assertEqual(data["{{宛名}}"], "山田 太郎 様")
        self.assertIn("A社に対し、令和8年9月1日付で", data["{{本文}}"])
        self.assertIn("令和8年10月1日付で", data["{{本文}}"])
        self.assertEqual(data["{{作成日}}"], "令和8年10月21日")
        self.assertIn("時効完成を認める旨", data["{{明細4}}"])
        self.assertEqual(data["{{署名1}}"], "大野法律事務所")
        self.assertIn("埼玉弁護士会・登録番号 12345", data["{{署名4}}"])
        self.assertFalse(any("{{" in v for v in data.values()))


class TestGenerateAndStore(_Base):
    def test_generate_docx_pdf_and_versions(self):
        with _approve_all():
            out1 = _run(C.generate(RID, 1, actor="human", today=datetime.date(2026, 10, 21)))
            out2 = _run(C.generate(RID, 1, actor="human", today=datetime.date(2026, 10, 21)))
        self.assertEqual((out1["version"], out2["version"]), (1, 2))              # 再生成で版が増える
        docs = _run(C.list_docs(RID))
        self.assertEqual([(d["slot"], d["version"]) for d in docs], [(1, 2), (1, 1)])
        d1 = _run(C.get_doc(RID, out1["doc_id"]))
        self.assertTrue(d1["pdf"].startswith(b"%PDF"))
        text = C._extract_text(d1["docx"])
        self.assertIn("時効援用完了報告書", text)
        self.assertIn("山田 太郎 様", text)
        self.assertIn("A社", text)
        for pii in (LINE_ID, PHONE, "09012345678", MAIL):                        # 相談者の連絡先は入らない
            self.assertNotIn(pii, text)
        self.assertEqual(d1["template_sha256"], C.TEMPLATE_SHA256)
        self.assertNotEqual(d1["doc_id"], out2["doc_id"])                        # 上書きなし

    def test_not_established_slot_refused_via_generate(self):
        with _approve_all():
            with self.assertRaises(C.GenerationRefused) as ctx:
                _run(C.generate(RID, 2, actor="human"))
        self.assertEqual(str(ctx.exception), "not_established")
        self.assertEqual(_run(C.list_docs(RID)), [])

    def test_pii_in_body_is_refused(self):
        # 回答要旨に相談者の電話番号が書かれていたら本文へ出さない（生成拒否）
        _run(F.record_result(RID, 1, result=F.RESULT_ESTABLISHED, answered_on=datetime.date(2026, 10, 1),
                             summary=f"折り返し先 {PHONE} に連絡済み"))
        with _approve_all():
            with self.assertRaises(C.GenerationRefused) as ctx:
                _run(C.generate(RID, 1, actor="human"))
        self.assertEqual(str(ctx.exception), "pii_in_body")
        self.assertEqual(_run(C.list_docs(RID)), [])

    def test_pdf_failure_stores_nothing(self):
        with _approve_all(), patch("contract_pdf.docx_to_pdf_bytes", side_effect=RuntimeError("boom")):
            with self.assertRaises(C.GenerationRefused) as ctx:
                _run(C.generate(RID, 1, actor="human"))
        self.assertEqual(str(ctx.exception), "pdf_failed")
        self.assertEqual(_run(C.list_docs(RID)), [])

    def test_unapproved_texts_refuse_generation(self):
        with self.assertRaises(C.GenerationRefused) as ctx:
            _run(C.generate(RID, 1, actor="human"))
        self.assertEqual(str(ctx.exception), "text_unapproved")


class TestView(_Base):
    def test_page_lists_established_only(self):
        r = _client.get(f"{V.PAGE}/{RID}", headers=_auth())
        self.assertEqual(r.status_code, 200)
        self.assertIn("枠 1: A社", r.text)
        self.assertNotIn("枠 2: B社", r.text)
        self.assertNotIn("山田", r.text)                                           # 顧客名は画面に出さない
        self.assertEqual(_client.get(f"{V.PAGE}/abc", headers=_auth()).status_code, 400)

    def test_generate_requires_confirm_then_creates_and_downloads(self):
        r = _client.post(f"{V.PAGE}/{RID}/generate/1", data={}, headers=_auth(), follow_redirects=False)
        self.assertEqual(r.status_code, 303)
        self.assertIn("done=confirm_required", r.headers["location"])
        self.assertEqual(_run(C.list_docs(RID)), [])
        with _approve_all():
            r = _client.post(f"{V.PAGE}/{RID}/generate/1", data={"confirm": "1"}, headers=_auth(), follow_redirects=False)
        self.assertEqual(r.status_code, 303)
        self.assertIn("done=generated", r.headers["location"])
        docs = _run(C.list_docs(RID))
        self.assertEqual(len(docs), 1)
        pdf = _client.get(f"{V.PAGE}/{RID}/doc/{docs[0]['doc_id']}.pdf", headers=_auth())
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertNotIn("山田", pdf.headers["content-disposition"])
        self.assertEqual(_client.get(f"{V.PAGE}/{RID}/doc/{'0' * 32}.pdf", headers=_auth()).status_code, 404)
        self.assertEqual(_client.get(f"{V.PAGE}/{RID}/doc/{docs[0]['doc_id']}.exe", headers=_auth()).status_code, 400)

    def test_refusal_reason_is_fixed_text(self):
        r = _client.post(f"{V.PAGE}/{RID}/generate/1", data={"confirm": "1"}, headers=_auth(), follow_redirects=False)
        self.assertIn("done=refused_text_unapproved", r.headers["location"])
        page = _client.get(r.headers["location"], headers=_auth())
        self.assertIn(V._NOTICES["refused_text_unapproved"], page.text)

    def test_unauthenticated_is_rejected(self):
        r = _client.get(f"{V.PAGE}/{RID}", follow_redirects=False)
        self.assertIn(r.status_code, (303, 401, 404))


if __name__ == "__main__":
    unittest.main()
