"""HOUKI-E2E-1（QUEUE-2026-10-13-E E-5）: 相続放棄 7 機能の通しテスト（fixture 1 案件・架空・kintone／LINE は全てモック）

統合ブランチ feature/houki-e2e-1（main + houki-dairinin-1 + houki-koseki-checklist-1 + houki-jijou-1 + houki-shokumu-link-1 +
houki-jukuryo-3）で、1 つの架空案件を次の順に通す:
  相談カード読取（AI 出力はモック・転記は実関数）→ ヒアリング転記（apply_hearing_fields）→ 弁護士入力の模擬（専権欄・status=受任・
  起算日_確定）→ 申述書 docx → 代理人目録 docx → 事情説明書 docx（下書きは固定文・素材一致検査は実関数）→ 戸籍チェックリスト初期行
  → 職務上請求接続（flag ON・提案封筒 1 件）→ 熟慮期間（期限・アラート候補・期限超過行）。
固定する規律（assert）:
- 相談者・業者への LINE 送信ゼロ（line_channel の HTTP 呼出 0・notify.push_line_message 0）。
- 承認なしの添付ゼロ（kintone.upload_file 0）・kintone 書込は「票で裁定された自動欄」の閉集合に収まる（status／弁護士専権欄をコードが
  書かない）。App 30 への起票は flag ON の提案封筒 1 件のみ。
- 各 docx はプレースホルダ残存なし・凍結検証を通過。
"""

import datetime
import os
import unittest
from unittest.mock import AsyncMock, patch

from test_houki_card_read import UID, V_ME, V_NAME, _FakeApp40, _filled, _tool_response  # noqa: E402

from hub import houki_card_read as r  # noqa: E402
from hub import houki_case_store as store  # noqa: E402
from hub import houki_dairinin as hd  # noqa: E402
from hub import houki_jijou as hj_jijou  # noqa: E402
from hub import houki_jukuryo as hj  # noqa: E402
from hub import houki_koseki_checklist as kc  # noqa: E402
from hub import houki_overdue as O  # noqa: E402
from hub import houki_shinjutsu as hs  # noqa: E402
from hub import houki_shokumu_link as L  # noqa: E402
from hub import kintone as hub_kintone  # noqa: E402
from hub import line_channel  # noqa: E402
from hub import notify as hub_notify  # noqa: E402
from hub import shokumu_plan  # noqa: E402
from test_houki_koseki_checklist import _props  # noqa: E402
from test_jikou_reply_q1a import _FakeClient, _reset_fake, _run  # noqa: E402

TODAY = datetime.date(2024, 6, 10)                 # 申述書・目録・上申書の作成日
LATER = datetime.date(2024, 9, 1)                  # 期限超過の確認日
OFFICE = {"OFFICE_ATTORNEY": "大野 太郎", "OFFICE_NAME": "大野法律事務所", "OFFICE_ZIP": "332-0031",
          "OFFICE_ADDRESS": "埼玉県川口市西青木2-1-45", "OFFICE_TEL": "048-299-2704",
          "OFFICE_FAX": "048-000-0000", "OFFICE_BAR": "埼玉弁護士会", "OFFICE_REG": "12345"}
# 弁護士が kintone 画面で入力する専権欄（コードは書かない＝本テストでは fixture に直接置く）
LAWYER_FIELDS = {
    "status": "受任", "管轄家庭裁判所": "さいたま家庭裁判所", "知った日の区分": "死亡の通知をうけた日",
    "相続の開始を知った日": "2024-05-03", "死亡日": "2024-05-01", "放棄の理由": "債務超過のため。",
    "起算日_確定": "2024-05-03", "起算点メモ": "親族からの連絡で死亡を知った日を起算日とする。",
}
# 票で裁定された「コードが自動で書いてよい欄」の閉集合（status・弁護士専権欄を含まない）
ALLOWED_AUTO_FIELDS = frozenset(store.HEARING_WRITABLE_FIELDS) | frozenset({
    "債権者一覧", r.FIELD_STATUS, "書類チェック"})
SECTIONS = {
    "kankei": "申述人は被相続人の子です。",
    "keii": "2024年5月3日に親族からの連絡で、被相続人の死亡（2024年5月1日）を知りました。",
    "riyu": "被相続人の債務を相続することを望まないためです。",
    "saimu": "消費者金融の借入があると認識しています。",
    "shobun": "遺産の処分は行っていません。",
}


def _docx_texts(data: bytes) -> list[str]:
    import io
    from docx import Document
    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out.extend(p.text for p in c.paragraphs)
    return out


class TestHoukiEndToEnd(unittest.TestCase):
    def setUp(self):
        r._generations.clear()
        self.addCleanup(r._generations.clear)
        _reset_fake()
        self.store = _FakeApp40()
        self.writes: list[dict] = []
        self.admin = AsyncMock(return_value=True)
        self.push = AsyncMock(return_value=True)
        self.upload = AsyncMock(return_value="fk")
        self.ai = AsyncMock(return_value=_tool_response(_filled(), name="read_card"))
        self.form = AsyncMock(return_value=_props())

        async def _update(app, record_id, fields, revision=None):
            self.writes.append(dict(fields))
            return await self.store.update_record(app, record_id, fields, revision)
        for p in (patch.object(hub_kintone, "search_records", self.store.search_records),
                  patch.object(hub_kintone, "create_record", self.store.create_record),
                  patch.object(hub_kintone, "get_record", self.store.get_record),
                  patch.object(hub_kintone, "update_record", _update),
                  patch.object(hub_kintone, "download_file", self.store.download_file),
                  patch.object(hub_kintone, "upload_file", self.upload),
                  patch.object(hub_kintone, "get_form_fields", self.form),
                  patch.object(r, "create_message_with_fallback", self.ai),
                  patch.object(hub_notify, "notify_admin_line", self.admin),
                  patch.object(hub_notify, "push_line_message", self.push),
                  patch.object(line_channel.httpx, "AsyncClient", _FakeClient),
                  patch.dict(os.environ, {**OFFICE, "APP_HOUKI": "40", "TOKEN_HOUKI": "t40"})):
            p.start()
            self.addCleanup(p.stop)
        kc.reset_schema_cache()
        self.addCleanup(kc.reset_schema_cache)

    def rec(self, rid: str) -> dict:
        return _run(self.store.get_record(None, rid))

    def lawyer_sets(self, rid: str, fields: dict) -> None:
        """人が kintone 画面で入力（revision が進む・コードの書込ではない）。"""
        case = self.store.cases[str(rid)]
        for k, v in fields.items():
            case[k] = {"value": v}
        case["$revision"] = {"value": str(int(case["$revision"]["value"]) + 1)}

    def test_seven_features_in_sequence(self):
        # ── 1. 相談カード読取（AI 出力モック → 実転記） ───────────────────────────
        rid = self.store.seed_card()
        self.assertEqual(_run(r.run_card_read(self.rec(rid))), "done")
        rec = self.rec(rid)
        self.assertEqual(store._v(rec, "被相続人氏名"), V_NAME)
        self.assertEqual(store._v(rec, "顧客名"), V_ME)
        self.assertEqual(store._v(rec, r.FIELD_STATUS), r.STATUS_DONE)
        self.assertEqual(self.store.creditors(rid), ["アコム"])
        self.assertEqual(store._v(rec, "status"), "問い合わせ")            # status はコードが動かさない

        # ── 2. ヒアリング転記（LINE の record_hearing と同じ関数・既存値は上書きしない） ──
        _rid, problems, choice = _run(store.apply_hearing_fields(UID, {
            "先順位者の放棄状況": "なし", "顧客名": "上書きされない名前"}, rec))
        self.assertEqual((_rid, problems, choice), (rid, [], []))
        rec = self.rec(rid)
        self.assertEqual(store._v(rec, "顧客名"), V_ME)                     # 追記専用（カードの値が残る）
        self.assertEqual(store._v(rec, "先順位者の放棄状況"), "なし")          # 空欄だけ埋まる
        self.assertEqual(store._v(rec, "被相続人本籍"), "不明")               # カード読取の値（不明）はそのまま

        # ── 3. 弁護士入力の模擬（専権欄・受任・起算日） ─────────────────────────
        self.lawyer_sets(rid, LAWYER_FIELDS)
        rec = self.rec(rid)
        self.assertTrue(hj.is_target(rec))

        # ── 4. 申述書 ───────────────────────────────────────────────────────────
        shinjutsu = hs.generate(rec, today=TODAY)
        self.assertTrue(shinjutsu.startswith(b"PK"))
        texts = _docx_texts(shinjutsu)
        self.assertNotIn("{{", "\n".join(texts))
        self.assertTrue(any(V_NAME in t for t in texts))

        # ── 5. 代理人目録 ───────────────────────────────────────────────────────
        dairinin = hd.generate(rec, today=TODAY)
        texts = _docx_texts(dairinin)
        self.assertNotIn("{{", "\n".join(texts))
        self.assertIn(f"被相続人　{V_NAME}　相続放棄申述事件", texts)
        self.assertIn("埼玉弁護士会", texts)

        # ── 6. 事情説明書（固定の下書き・素材一致検査は実関数） ──────────────────
        self.assertEqual(hj_jijou.check_against_source(SECTIONS, rec), [])
        jijou = hj_jijou.generate(rec, SECTIONS, today=TODAY)
        texts = _docx_texts(jijou)
        self.assertNotIn("{{", "\n".join(texts))
        self.assertTrue(any(SECTIONS["shobun"] in t for t in texts))
        with self.assertRaises(hj_jijou.JijouGuardError):                  # 素材に無い金額は生成しない
            hj_jijou.generate(rec, dict(SECTIONS, saimu="債務は 300万円 です。"), today=TODAY)

        # ── 7. 戸籍チェックリスト初期行（受任時・1 回だけ） ──────────────────────
        res = _run(kc.init_checklist(rid, rec))
        self.assertIn(res, (kc.RESULT_CREATED, kc.RESULT_INDIVIDUAL))
        rec = self.rec(rid)
        rows = rec["書類チェック"]["value"]
        self.assertGreaterEqual(len(rows), len(kc.COMMON_ROWS))
        names = [row["value"][kc.COL_NAME]["value"] for row in rows]
        self.assertIn("被相続人住民票除票", names)
        self.assertEqual(_run(kc.init_checklist(rid, rec)), kc.RESULT_EXISTS)   # 冪等

        # ── 8. 職務上請求接続（flag ON・提案封筒 1 件・冪等） ─────────────────────
        created: list = []
        search = AsyncMock(return_value=[])

        async def _create(app, fields):
            created.append((app, fields))
            search.return_value = [{"$id": {"value": "301"}, "チャネル固有データ": {"value": fields["チャネル固有データ"]},
                                    "発送ステータス": {"value": fields["発送ステータス"]}, "実行済み": {"value": "no"}}]
            return "301"
        app31 = AsyncMock(return_value={"municipality": "木曽郡上松町", "app31_record_id": "7", "有効": "yes", "fallback": ""})
        with patch.object(hub_kintone, "search_records", search), patch.object(hub_kintone, "create_record", AsyncMock(side_effect=_create)), \
                patch.object(shokumu_plan, "resolve_app31", app31), \
                patch.dict(os.environ, {"SHOKUMU_PLAN_ENABLED": "1", "APP_SHIPPING": "30", "TOKEN_SHIPPING": "t30"}):
            out = _run(L.process_record(rec))
            self.assertEqual(out["status"], L.RESULT_FILED)
            self.assertEqual(len(created), 1)
            fields = created[0][1]
            self.assertEqual((fields["ユニット種別"], fields["実行済み"], fields["案件レコードID"]), ("相続放棄", "no", rid))
            self.assertEqual(fields["発送ステータス"], shokumu_plan.STATUS_PENDING)
            self.assertEqual(_run(L.process_record(rec))["status"], L.RESULT_ALREADY)   # 冪等（open な同一キー）
            self.assertEqual(len(created), 1)
        with patch.dict(os.environ, {"SHOKUMU_PLAN_ENABLED": ""}):
            self.assertEqual(_run(L.process_record(rec))["status"], L.RESULT_DISABLED)

        # ── 9. 熟慮期間（期限・アラート候補・期限超過） ─────────────────────────
        comp = hj.compute(rec, TODAY)
        self.assertEqual((comp.start, comp.deadline), (datetime.date(2024, 5, 3), datetime.date(2024, 8, 2)))
        self.assertEqual(comp.remaining, (comp.deadline - TODAY).days)
        self.assertEqual(hj.alert_candidates(comp.remaining), [])
        self.assertEqual(hj.alert_candidates(10), [14])
        self.assertEqual(hj.alert_candidates(3), [14, 7])
        over = O.overdue_rows([rec], LATER)
        self.assertEqual([(x.record_id, x.deadline, x.remaining) for x in over], [(rid, datetime.date(2024, 8, 2), -30)])
        self.assertEqual(O.overdue_rows([rec], TODAY), [])
        self.assertIn("1 件", O.notice_text(LATER, 1))
        self.assertNotIn(V_ME, O.notice_text(LATER, 1))

        # ── 規律: 送信ゼロ・承認なし添付ゼロ・自動書込は閉集合内 ────────────────
        self.assertEqual(_FakeClient.calls, [])
        self.assertEqual(self.push.call_count, 0)
        self.assertEqual(self.upload.call_count, 0)
        written = set()
        for w in self.writes:
            written |= set(w.keys())
        self.assertTrue(written <= ALLOWED_AUTO_FIELDS, written - ALLOWED_AUTO_FIELDS)
        self.assertNotIn("status", written)
        for k in LAWYER_FIELDS:
            self.assertNotIn(k, written, k)
        for call in self.admin.call_args_list:                               # 管理者通知に値を載せない
            text = str(call.args[0] if call.args else call.kwargs.get("text", ""))
            self.assertNotIn(V_ME, text)
            self.assertNotIn(V_NAME, text)


if __name__ == "__main__":
    unittest.main()
