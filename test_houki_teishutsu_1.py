"""HOUKI-TEISHUTSU-1（QUEUE-2026-10-16-R F-3）: 相続放棄 家裁提出パッケージ（提出前チェック・送付状・費用定数・提出日の記録）

固定する仕様:
- 1) 提出前チェック: 申述書・代理人目録・事情説明書の添付、戸籍チェックリストの充足（戸籍系の行が全て 受領／不要・1 行以上 受領）、
  申述人情報の必須欄（hub/houki_shinjutsu の拒否理由を流用）、熟慮期間の残日数。不足が 1 つでもあれば画面に「送付状を生成」
  「提出日を記録」のフォームを出さない（POST しても not_ready）。
- 2) 送付状 docx: 雛形 docx_templates/houki/送付状.docx の SHA pin・全文 MS Mincho・太字なし・表／テキストボックスなし・各 {{ }} 単一 run。
  同封書類一覧は実際に添付されている生成済み書類＋受領済み戸籍等＋印紙・郵券。FILE 欄「送付状」があれば添付、無ければダウンロード。
- 3) 収入印紙・予納郵券は hub/houki_fees の定数（要大野確認）。未登録の家裁は「要確認」。本票の他ファイルに金額を直書きしない。
- 4) 提出日の記録は App 40「申述提出日」に $revision CAS（人の操作・明示確認）。記録後は既存 HOUKI-JUKURYO の対象外＝通知が出ない。
- pin: 新 module の sink 違反ゼロ・allowlist 追加ゼロ・main は末尾結線のみ・送信なし・書く欄は 申述提出日／送付状 だけ。
"""

import asyncio
import datetime
import hashlib
import io
import os
import re
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document

for _k, _v in {
    "KINTONE_SUBDOMAIN": "testsub", "LINE_CHANNEL_SECRET": "dummy_secret",
    "LINE_CHANNEL_ACCESS_TOKEN": "dummy_token", "ANTHROPIC_API_KEY": "dummy_key",
    "KINTONE_APP_ID": "21", "KINTONE_API_TOKEN": "dummy",
    "SOUZOKU_KINTONE_APP_ID": "26", "SOUZOKU_KINTONE_API_TOKEN": "dummy",
    "CLOUDSIGN_CLIENT_ID": "c", "CLOUDSIGN_WEBHOOK_SECRET": "cs",
    "KINTONE_WEBHOOK_TOKEN": "kintone-token", "DOCUMENT_WEBHOOK_SECRET": "d",
    "APP_APPROVAL": "29", "TOKEN_APPROVAL": "d", "HEALTHCHECK_DISABLED": "1",
    "STRIPE_WEBHOOK_SECRET": "w", "GOOGLE_VISION_API_KEY": "dummy_vision",
    "APP_HOUKI": "40", "TOKEN_HOUKI": "t40",
}.items():
    os.environ.setdefault(_k, _v)

from fastapi.testclient import TestClient  # noqa: E402

import main  # noqa: E402
from hub import db  # noqa: E402
from hub import houki_fees as FEES  # noqa: E402
from hub import houki_jukuryo as hj  # noqa: E402
from hub import houki_koseki_checklist as K  # noqa: E402
from hub import houki_teishutsu as T  # noqa: E402
from hub import kintone  # noqa: E402
from hub import send_ledger as sl  # noqa: E402
from hub import webapp_houki_submit_view as V  # noqa: E402
from hub.webapp_auth import MIN_ITERATIONS, hash_password, issue_session  # noqa: E402
from test_sink_ast_policy import scan_source  # noqa: E402

REPO = Path(__file__).parent
_client = TestClient(main.app)
_ENV = {"WEBAPP_PASSWORD_HASH": hash_password("pw", iterations=MIN_ITERATIONS), "WEBAPP_SESSION_SECRET": "s" * 32,
        "OFFICE_ATTORNEY": "大野 太郎", "OFFICE_NAME": "大野法律事務所", "OFFICE_ADDRESS": "埼玉県川口市西青木2-1-45",
        "OFFICE_TEL": "048-299-2704", "OFFICE_FAX": "048-299-2705"}
USER = "U" + "a" * 32
TODAY = datetime.date(2026, 10, 16)
ADMIN = "U" + "c" * 32


def _row(name, target="", state=K.STATE_RECEIVED):
    return {"value": {K.COL_NAME: {"value": name}, K.COL_TARGET: {"value": target},
                      K.COL_METHOD: {"value": "職務上請求"}, K.COL_STATE: {"value": state}}}


ROWS_OK = [_row(K.DOC_JOHYO, "被相続人"), _row(K.DOC_POA, "申述人"), _row(K.DOC_STAMP, "事務所", K.STATE_NOT_NEEDED),
           _row(K.DOC_POSTAGE, "事務所", K.STATE_NOT_NEEDED), _row(K.DOC_DEATH_KOSEKI, "被相続人"),
           _row(K.DOC_APPLICANT_KOSEKI, "申述人")]


def _rec(**over):
    base = {"$id": "40", "$revision": "3", "status": "申述書作成", "顧客名": "山田 花子", "住所": "埼玉県○○市1-2-3",
            "生年月日": "1980-01-02", "電話番号": "090-1111-2222", "LINEユーザーID": USER, "被相続人氏名": "山田 太郎",
            "被相続人ふりがな": "やまだ たろう", "被相続人本籍": "埼玉県", "被相続人最後の住所": "埼玉県", "死亡日_申告": "2026-08-10",
            "相続の開始を知った日": "2026-08-20", "知った日の区分": "死亡の通知をうけた日", "続柄": "子", "放棄の理由": "債務超過のため。",
            "管轄家庭裁判所": "さいたま家庭裁判所", "起算日_確定": "2026-08-20", "申述提出日": "",
            "申述書": [{"fileKey": "a"}], "代理人目録": [{"fileKey": "b"}], "事情説明書": [{"fileKey": "c"}], "書類チェック": ROWS_OK}
    base.update(over)
    return {k: {"value": v} for k, v in base.items()}


def _auth():
    return {"Cookie": f"webapp_session={issue_session()}"}


class TestFees(unittest.TestCase):
    def test_constants_and_lookup(self):
        self.assertEqual(FEES.stamp_yen(1), FEES.STAMP_YEN_PER_APPLICANT)
        self.assertEqual(FEES.stamp_yen(3), FEES.STAMP_YEN_PER_APPLICANT * 3)
        self.assertEqual(FEES.STAMP_STATUS, "要大野確認")
        self.assertTrue(all(p.status == "要大野確認" for p in FEES.POSTAGE_TABLE))
        p = FEES.postage_for("さいたま家庭裁判所")
        self.assertEqual(p.court, "さいたま家庭裁判所")
        u = FEES.postage_for("札幌家庭裁判所")
        self.assertIsNone(u.yen)
        self.assertIn("要確認", u.breakdown)
        self.assertEqual(FEES.fmt_yen(None), "要確認")
        self.assertEqual(FEES.fmt_yen(1200), "1,200円")
        self.assertIn("要大野確認", (REPO / "hub" / "houki_fees.py").read_text(encoding="utf-8"))


class TestChecklist(unittest.TestCase):
    def test_ready_when_everything_present(self):
        items = T.checklist(_rec(), TODAY)
        self.assertEqual([i.key for i in items], ["申述書", "代理人目録", "事情説明書", "koseki", "required", "jukuryo"])
        self.assertTrue(T.is_ready(items), [(i.key, i.detail) for i in items if not i.ok])
        self.assertIn("期限 2026-11-19", items[-1].detail)                 # 8/20 + 3 か月の応当日の前日
        self.assertEqual(T.submitted_on(_rec()), "")

    def test_each_shortage_blocks(self):
        for over, key in ((dict(事情説明書=[]), "事情説明書"), (dict(代理人目録=[]), "代理人目録"), (dict(申述書=[]), "申述書"),
                          (dict(書類チェック=[]), "koseki"), (dict(書類チェック=ROWS_OK[:5] + [_row(K.DOC_APPLICANT_KOSEKI, "申述人", "請求中")]), "koseki"),
                          (dict(住所=""), "required"), (dict(顧客名=""), "required"), (dict(起算日_確定=""), "jukuryo"),
                          (dict(起算日_確定="2026-05-01"), "jukuryo")):
            items = T.checklist(_rec(**over), TODAY)
            bad = [i.key for i in items if not i.ok]
            self.assertEqual(bad, [key], (over, bad))
            self.assertFalse(T.is_ready(items))
        it = next(i for i in T.checklist(_rec(住所=""), TODAY) if i.key == "required")
        self.assertIn("未入力: 住所", it.detail)
        it = next(i for i in T.checklist(_rec(起算日_確定="2026-05-01"), TODAY) if i.key == "jukuryo")
        self.assertIn("期限超過", it.detail)

    def test_enclosures_follow_attachments_and_received_rows(self):
        enc = T.enclosures(_rec())
        self.assertEqual(enc[:3], ["相続放棄申述書", "代理人目録", "上申書（事情説明書）"])
        self.assertIn("被相続人住民票除票（被相続人）", enc)
        self.assertIn("委任状", enc)
        self.assertIn(f"収入印紙　{FEES.fmt_yen(FEES.stamp_yen(1))}分", enc)
        self.assertIn("予納郵券　要確認", enc)
        enc2 = T.enclosures(_rec(事情説明書=[], 書類チェック=[_row(K.DOC_DEATH_KOSEKI, "被相続人", "請求中")]))
        self.assertNotIn("上申書（事情説明書）", enc2)
        self.assertFalse(any("戸籍" in x for x in enc2))                  # 未受領は同封しない
        txt = T.enclosure_text(["A", "B"])
        self.assertEqual(txt, "１　A\n２　B")


class TestTemplate(unittest.TestCase):
    def test_template_pin_structure_and_fonts(self):
        data = (REPO / T.TEMPLATE_PATH).read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(), T.TEMPLATE_SHA256)
        import importlib.util
        spec = importlib.util.spec_from_file_location("mk", REPO / "scripts" / "make_soufujo_template.py")
        mk = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mk)
        self.assertEqual(hashlib.sha256(mk.build()).hexdigest(), T.TEMPLATE_SHA256)      # 決定的・再現可能
        doc = Document(io.BytesIO(data))
        T.verify_structure(doc)
        self.assertEqual(doc.tables, [])
        xml = doc.element.xml
        self.assertNotIn("txbx", xml)
        self.assertNotIn("<w:b/>", xml)
        self.assertEqual(doc.styles["Normal"].font.name, T.FONT_JA)
        for p in doc.paragraphs:
            for r in p.runs:
                self.assertFalse(r.bold)

    def test_tampered_template_rejected(self):
        data = (REPO / T.TEMPLATE_PATH).read_bytes()
        doc = Document(io.BytesIO(data))
        doc.paragraphs[8].runs[0].bold = True
        out = io.BytesIO()
        doc.save(out)
        with self.assertRaises(T.TeishutsuIntegrityError):
            T.verify_structure(Document(io.BytesIO(out.getvalue())))
        with self.assertRaises(T.TeishutsuIntegrityError):
            T.verify_template_bytes(out.getvalue())
        with patch.object(T, "TEMPLATE_SHA256", "0" * 64):
            with self.assertRaises(T.TeishutsuIntegrityError):
                T.generate(_rec(), TODAY)

    def test_generate_fills_and_matches_enclosures(self):
        with patch.dict(os.environ, _ENV):
            out = T.generate(_rec(), TODAY)
            fill = T.build_fill(_rec(), TODAY)
        doc = Document(io.BytesIO(out))
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("さいたま家庭裁判所　御中", texts)
        self.assertIn("事件名　被相続人　山田 太郎　相続放棄申述事件", texts)
        self.assertIn("申述人　山田 花子", texts)
        self.assertIn("令和8年10月16日", texts)
        self.assertIn("弁護士　大野 太郎", texts)
        body = "\n".join(texts)
        for enc in T.enclosures(_rec()):
            self.assertIn(enc, body)                                          # 同封一覧＝実際の添付・受領
        self.assertNotIn("{{", body)
        self.assertNotIn(USER, body)
        self.assertNotIn("090-1111-2222", body)
        T.verify_structure(doc, filled=fill)
        with self.assertRaises(T.TeishutsuRejection):
            T.generate(_rec(管轄家庭裁判所=""), TODAY)


class TestView(unittest.TestCase):
    def setUp(self):
        self._env = patch.dict(os.environ, _ENV)
        self._env.start()
        V.clear_caches()
        self.record = _rec()
        self.get_record = AsyncMock(side_effect=lambda app, rid: self.record)
        self.upload = AsyncMock(return_value="fk1")
        self.update = AsyncMock()
        self.fields = AsyncMock(return_value={})                                 # 既定: 「送付状」欄なし
        self._p = [patch.object(kintone, "get_record", self.get_record), patch.object(kintone, "upload_file", self.upload),
                   patch.object(kintone, "update_record", self.update), patch.object(kintone, "get_form_fields", self.fields),
                   patch.object(T, "today_jst", lambda: TODAY)]
        for p in self._p:
            p.start()

    def tearDown(self):
        for p in self._p:
            p.stop()
        V.clear_caches()
        self._env.stop()

    def _form(self):
        r = _client.get(f"{V.PAGE}/40", headers=_auth())
        self.assertEqual(r.status_code, 200)
        op = r.text.split('name="op_id" value="')[1].split('"')[0]
        tk = r.text.split('name="ticket" value="')[1].split('"')[0]
        return r, {"op_id": op, "ticket": tk}

    def test_gate_and_shortage_hides_buttons(self):
        self.assertEqual(_client.get(f"{V.PAGE}/40", follow_redirects=False).status_code, 303)
        self.assertEqual(_client.get(f"{V.PAGE}/abc", headers=_auth()).status_code, 400)
        self.record = _rec(事情説明書=[])
        r = _client.get(f"{V.PAGE}/40", headers=_auth())
        self.assertEqual(r.status_code, 200)
        self.assertIn("不足", r.text)
        self.assertNotIn("送付状を生成", r.text)
        self.assertNotIn("提出日を記録", r.text)
        self.assertNotIn(f"{V.PAGE}/40/generate", r.text)
        self.assertNotIn(USER, r.text)
        self.assertNotIn("090-1111-2222", r.text)
        self.assertNotIn("埼玉県○○市", r.text)
        self.upload.assert_not_called()
        self.update.assert_not_called()
        j = _client.get(f"{V.API}/40", headers=_auth()).json()
        self.assertFalse(j["ready"])
        self.assertEqual(j["fees"]["stamp_yen"], FEES.stamp_yen(1))

    def test_ready_shows_both_forms_and_fees(self):
        r, _ = self._form()
        self.assertIn("送付状を生成", r.text)
        self.assertIn("提出日を記録", r.text)
        self.assertIn(FEES.fmt_yen(FEES.stamp_yen(1)), r.text)
        self.assertIn("要大野確認", r.text)
        self.assertIn("要確認", r.text)                                           # 郵券は未登録（要確認）
        self.assertTrue(_client.get(f"{V.API}/40", headers=_auth()).json()["ready"])

    def test_generate_download_when_no_field_and_attach_when_field(self):
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/generate", data=f, headers=_auth(), follow_redirects=False)
        self.assertIn("confirm_required", r.headers["location"])
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1", ticket="bad"), headers=_auth(), follow_redirects=False)
        self.assertEqual(r.status_code, 400)
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertEqual(r.status_code, 200)                                      # 欄なし → ダウンロード
        self.assertIn("wordprocessingml", r.headers["content-type"])
        self.assertIn("soufujo_40.docx", r.headers["content-disposition"])
        texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
        self.assertIn("申述人　山田 花子", texts)
        self.upload.assert_not_called()
        V.clear_caches()
        self.fields.return_value = {T.FIELD_SOUFUJO: {"type": "FILE"}}
        r = _client.post(f"{V.PAGE}/40/generate", data=dict(f, confirm="1"), headers=_auth(), follow_redirects=False)
        self.assertIn("done=generated", r.headers["location"])
        self.assertEqual(self.upload.await_count, 1)
        self.assertEqual(self.update.await_args.args[2], {T.FIELD_SOUFUJO: [{"fileKey": "fk1"}]})
        self.assertEqual(self.update.await_args.kwargs["revision"], "3")

    def test_not_ready_and_stale_on_post(self):
        _, f = self._form()
        self.record = _rec(事情説明書=[])
        r = _client.post(f"{V.PAGE}/40/submitted", data=dict(f, confirm="1", submitted_on="2026-10-16"), headers=_auth(),
                         follow_redirects=False)
        self.assertIn("done=not_ready", r.headers["location"])
        self.record = _rec()
        self.record["$revision"] = {"value": "4"}
        r = _client.post(f"{V.PAGE}/40/submitted", data=dict(f, confirm="1", submitted_on="2026-10-16"), headers=_auth(),
                         follow_redirects=False)
        self.assertIn("done=stale", r.headers["location"])
        self.update.assert_not_called()

    def test_record_submitted_writes_only_that_field_then_hides_buttons(self):
        _, f = self._form()
        r = _client.post(f"{V.PAGE}/40/submitted", data=dict(f, confirm="1", submitted_on="2026/10/16"), headers=_auth(),
                         follow_redirects=False)
        self.assertIn("done=bad_input", r.headers["location"])
        r = _client.post(f"{V.PAGE}/40/submitted", data=dict(f, confirm="1", submitted_on="2026-10-16"), headers=_auth(),
                         follow_redirects=False)
        self.assertIn("done=submitted", r.headers["location"])
        self.assertEqual(self.update.await_args.args[1:3], ("40", {T.FIELD_SUBMITTED: "2026-10-16"}))
        self.assertEqual(self.update.await_args.kwargs["revision"], "3")
        self.record = _rec(申述提出日="2026-10-16")
        r = _client.get(f"{V.PAGE}/40", headers=_auth())
        self.assertIn("提出済み", r.text)
        self.assertNotIn("提出日を記録", r.text)
        r = _client.post(f"{V.PAGE}/40/submitted", data=dict(f, confirm="1", submitted_on="2026-10-16"), headers=_auth(),
                         follow_redirects=False)
        self.assertIn("done=already_submitted", r.headers["location"])
        self.assertEqual(self.update.await_count, 1)


class TestJukuryoStops(unittest.TestCase):
    """提出日記録後は既存 HOUKI-JUKURYO の対象外（検索条件・is_target・日次実走で通知ゼロ）。"""

    def setUp(self):
        self._dir = tempfile.mkdtemp(prefix="teishutsu_")
        self._env = patch.dict(os.environ, {"DATABASE_URL": f"sqlite+aiosqlite:///{self._dir}/n.db",
                                            "DISPATCHBOT_CHANNEL_ACCESS_TOKEN": "biz-token", **_ENV})
        self._env.start()
        db.reset_for_tests()

        async def _create():
            async with db.get_async_engine().begin() as c:
                await c.run_sync(sl.metadata.create_all)
        asyncio.run(_create())
        db.reset_for_tests()
        self.push = AsyncMock(return_value=True)
        self.update = AsyncMock()
        self._p = [patch.object(hj, "_push_admin_http", self.push), patch.object(hj, "get_admin_line_user_id", lambda: ADMIN),
                   patch.object(hj, "business_channel_allowlist", lambda: frozenset({ADMIN})),
                   patch.object(hj, "_today_jst", lambda: datetime.date(2026, 11, 10)),     # 期限 11/19 の 9 日前（7 日前通知の手前・14 日前は到来）
                   patch.object(kintone, "update_record", self.update), patch.object(kintone, "get_record", AsyncMock())]
        for p in self._p:
            p.start()

    def tearDown(self):
        for p in self._p:
            p.stop()
        db.reset_for_tests()
        self._env.stop()
        shutil.rmtree(self._dir, ignore_errors=True)

    def _run_daily(self, records):
        with patch.object(hj, "fetch_all_targets", AsyncMock(return_value=records)):
            try:
                return asyncio.run(hj.jukuryo_daily_check())
            finally:
                db.reset_for_tests()

    def test_submitted_record_is_not_a_target(self):
        rec_open = _rec()
        rec_done = _rec(申述提出日="2026-10-16")
        self.assertTrue(hj.is_target(rec_open))
        self.assertFalse(hj.is_target(rec_done))
        self.assertIn(f'{hj.FIELD_SUBMITTED} = ""', hj.search_query())
        self._run_daily([rec_open])
        self.assertEqual(self.push.await_count, 1)                                # 未提出: 14 日前通知が 1 通
        self.push.reset_mock()
        self._run_daily([rec_done])                                               # 提出済み（検索から外れていなくても）: 通知ゼロ
        self.assertEqual(self.push.await_count, 0)


class TestPins(unittest.TestCase):
    def test_sink_allowlist_wiring_and_write_set(self):
        allow = (REPO / "redaction_sink_allowlist.json").read_text(encoding="utf-8")
        for name in ("hub/houki_teishutsu.py", "hub/houki_fees.py", "hub/webapp_houki_submit_view.py", "scripts/make_soufujo_template.py"):
            self.assertEqual(scan_source((REPO / name).read_text(encoding="utf-8"), name), [], name)
        for mod in ("houki_teishutsu", "houki_fees", "webapp_houki_submit_view", "make_soufujo_template"):
            self.assertNotIn(mod, allow)
        src = (REPO / "main.py").read_text(encoding="utf-8").split("\n")
        self.assertIn("include_before_catch_all(app, houki_submit_router)", "\n".join(src[-3:]))
        for name in ("hub/houki_teishutsu.py", "hub/webapp_houki_submit_view.py"):
            body = (REPO / name).read_text(encoding="utf-8").split('"""', 2)[2]
            for bad in ("line_channel", "push_text", "notify_admin_line", "httpx", "delete_record"):
                self.assertNotIn(bad, body, name)
        vsrc = (REPO / "hub" / "webapp_houki_submit_view.py").read_text(encoding="utf-8")
        writes = re.findall(r"update_record\(APP_HOUKI_CASE, rid, \{([^}]+)\}", vsrc)
        self.assertEqual(sorted(w.split(":")[0].strip() for w in writes), ["T.FIELD_SOUFUJO", "T.FIELD_SUBMITTED"])
        tsrc = (REPO / "hub" / "houki_teishutsu.py").read_text(encoding="utf-8")
        self.assertNotIn("800", tsrc)                                             # 金額は hub/houki_fees のみ
        self.assertNotIn("800", vsrc)
        for r in main.app.router.routes:
            if getattr(r, "path", "").startswith((V.PAGE, V.API)):
                self.assertTrue(getattr(r.endpoint, "__webapp_gate__", False), r.path)
        self.assertEqual(T.FIELD_SUBMITTED, hj.FIELD_SUBMITTED)


if __name__ == "__main__":
    unittest.main()
